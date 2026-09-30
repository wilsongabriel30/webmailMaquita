# -*- coding: utf-8 -*-
"""Prueba de vínculos con filtro y apilados, en la caja de arena de Wilson.

Crea un libro de destino en /PRUEBAS FORMULARIOS/ASC, un vínculo con filtro
(O = ASC - Comunicación) sobre el consolidado de la caja de arena y otro
apilado debajo (O = ASC - Asesorías); refresca; comprueba filas; cambia el
filtro a algo más corto y comprueba que lo sobrante se limpia. Al final borra
lo creado."""
import io
import os
import sys
import warnings

warnings.simplefilter('ignore')
sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
import openpyxl
import almacen_bd as bd
import nucleo_archivos as nucleo
import api_vinculos
import vinculos_filtro
from seguridad_rutas import ruta_fisica

USUARIO = 14
CARPETA = '/PRUEBAS FORMULARIOS/ASC'
DESTINO = CARPETA + '/prueba-vinculo-filtro.xlsx'
ORIGEN = CARPETA + '/Mapa de inversiones/Mapa general de inversiones 2026.xlsx'
fallos = []


def comprobar(c, texto, detalle=''):
    if not c:
        fallos.append(texto)
    print(('OK   ' if c else 'MAL  ') + texto + ((' — ' + str(detalle)) if (detalle and not c) else ''))


def filas_destino(hoja='Datos', celda='A2'):
    wb = openpyxl.load_workbook(ruta_fisica(USUARIO, DESTINO), read_only=True)
    ws = wb[hoja]
    n = 0
    col = openpyxl.utils.cell.coordinate_to_tuple(celda)[1]
    fila0 = openpyxl.utils.cell.coordinate_to_tuple(celda)[0]
    valores = []
    for f in ws.iter_rows(min_row=fila0, min_col=col, max_col=col + 7, values_only=True):
        if any(v not in (None, '') for v in f):
            n += 1
            valores.append(f)
    wb.close()
    return n, valores


ids = []
try:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Datos'
    ws['A1'] = 'PRUEBA vínculo con filtro'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nucleo.subir(USUARIO, CARPETA, 'prueba-vinculo-filtro.xlsx', buf)
    comprobar(os.path.exists(ruta_fisica(USUARIO, DESTINO)), 'libro de destino creado')

    vinculos_filtro.asegurar_columnas()
    cab = bd.ejecutar(
        "INSERT INTO vinculos_datos (origen_usuario, origen_ruta, origen_hoja, origen_rango, "
        "destino_usuario, destino_ruta, destino_hoja, destino_celda, creado_por, filtro_columna, filtro_valores) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING *",
        (USUARIO, ORIGEN, 'ConsolidadoP', 'H4:CN8000', USUARIO, DESTINO, 'Datos', 'A2', USUARIO, 'O', 'ASC - Comunicación'))
    ids.append(cab['id'])
    ok, msg = api_vinculos._refrescar(dict(cab))
    comprobar(ok, 'refresco del vínculo con filtro', msg)
    n, vals = filas_destino()
    comprobar(n == 124, 'trae solo las filas de Comunicación (124)', n)
    comprobar(all(str(v[7]).strip() == 'ASC - Comunicación' for v in vals), 'todas las filas cumplen el filtro')

    mie = bd.ejecutar(
        "INSERT INTO vinculos_datos (origen_usuario, origen_ruta, origen_hoja, origen_rango, "
        "destino_usuario, destino_ruta, destino_hoja, destino_celda, creado_por, filtro_columna, filtro_valores, apilar_id) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING *",
        (USUARIO, ORIGEN, 'ConsolidadoP', 'H4:CN8000', USUARIO, DESTINO, 'Datos', 'A2', USUARIO, 'Col8', 'asc - asesorias', cab['id']))
    ids.append(mie['id'])
    ok, msg = api_vinculos._refrescar(dict(mie))
    comprobar(ok, 'refresco desde el vínculo apilado (escribe la pila entera)', msg)
    n, vals = filas_destino()
    comprobar(n == 124 + 42, 'apilado: Comunicación + Asesorías (166), con «Col8» y sin tildes', n)
    comprobar(str(vals[124][7]).strip() == 'ASC - Asesorías', 'las de Asesorías van debajo')
    comprobar(api_vinculos._destino_al_dia(dict(cab)), 'el destino se reconoce al día (no se reescribe en balde)')

    bd.ejecutar("UPDATE vinculos_datos SET filtro_valores = %s WHERE id = %s", ('ASC - Eventos Nacionales', cab['id']))
    cab2 = dict(bd.consultar('SELECT * FROM vinculos_datos WHERE id = %s', (cab['id'],))[0])
    ok, msg = api_vinculos._refrescar(cab2)
    comprobar(ok, 'refresco tras cambiar el filtro a Eventos Nacionales', msg)
    n, vals = filas_destino()
    comprobar(n == 43 + 42, 'lo sobrante se limpió (43 + 42 = 85)', n)
    comprobar(filas_destino('Datos', 'A1')[1][0][0] == 'PRUEBA vínculo con filtro', 'lo que hay por encima no se toca')

    # Lo que devuelve la API en vivo es la misma pila
    from api_vinculos_vivo import _recortar
    vivo, _ = vinculos_filtro.matriz_de(cab2, api_vinculos._leer_rango)
    comprobar(len(_recortar(vivo)) == 85, 'la vista en vivo entrega la pila completa (85)', len(vivo))
    comprobar(vinculos_filtro.indice_columna('O', 'H4:CN8000') == 7 and vinculos_filtro.indice_columna('Col8', 'H4:CN8000') == 7,
              '«O» y «Col8» señalan la misma columna del rango')
finally:
    for i in ids:
        bd.ejecutar('DELETE FROM vinculos_datos WHERE id = %s', (i,))
    try:
        nucleo.enviar_a_papelera(USUARIO, DESTINO)
    except Exception as e:
        try:
            os.unlink(ruta_fisica(USUARIO, DESTINO))
        except Exception:
            pass
        print('AVISO al borrar el libro de prueba:', e)

print('\n%s — %d fallos' % ('TODO BIEN' if not fallos else 'HAY FALLOS', len(fallos)))
for f in fallos:
    print('  · ' + f)
sys.exit(1 if fallos else 0)
