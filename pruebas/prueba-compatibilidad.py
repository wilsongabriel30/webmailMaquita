# -*- coding: utf-8 -*-
"""Prueba de compatibilidad_xlsx: lo que sale del editor tiene que abrirse con
listas y colores en un programa que NO entiende la extensión de Microsoft.
openpyxl es justamente uno de ellos, así que sirve de juez.

Los dos archivos de muestra están al lado: `fixture-oo-listas.xlsx` es una
salida REAL del Document Server (lista con colores en E9, todo en la extensión)
y `fixture-normal.xlsx` una hoja corriente sin nada que traducir.

Uso: cd /home/sistemas/almacen-maquita/pruebas && python3 prueba-compatibilidad.py
"""
import os
import sys
import warnings
import xml.etree.ElementTree as ET
import zipfile

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
import compatibilidad_xlsx as compat                              # noqa: E402
from openpyxl import load_workbook                                # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
ORIGEN = sys.argv[1] if len(sys.argv) > 1 else AQUI + '/fixture-oo-listas.xlsx'
SALIDA = '/tmp/compat-salida.xlsx'

rojo = verde = 0


def bien(ok, que):
    global rojo, verde
    if ok:
        verde += 1
    else:
        rojo += 1
    print(('OK   ' if ok else 'MAL  ') + que)


# ── antes: openpyxl no ve nada (así se ve fuera del Drive) ───────────────
with warnings.catch_warnings():
    warnings.simplefilter('ignore')
    h = load_workbook(ORIGEN).worksheets[0]
antes_dv = len(h.data_validations.dataValidation)
antes_cf = len(list(h.conditional_formatting))
print('antes de la pasada: validaciones=%d colores=%d' % (antes_dv, antes_cf))

# ── la pasada ────────────────────────────────────────────────────────────
bien(compat.convertir(ORIGEN, SALIDA), 'la pasada dice que tradujo algo')

with zipfile.ZipFile(SALIDA) as z:
    hoja = z.read('xl/worksheets/sheet1.xml').decode()
    estilos = z.read('xl/styles.xml').decode()
    nombres_salida = set(z.namelist())
with zipfile.ZipFile(ORIGEN) as z:
    nombres_origen = set(z.namelist())

bien(nombres_salida == nombres_origen, 'no se pierde ni se añade ninguna pieza del ZIP')
for texto, quien in ((hoja, 'la hoja'), (estilos, 'styles.xml')):
    try:
        ET.fromstring(texto)
        bien(True, '%s sigue siendo XML válido' % quien)
    except ET.ParseError as e:
        bien(False, '%s sigue siendo XML válido (%s)' % (quien, e))

bien('<dataValidation ' in hoja, 'la lista queda en la forma clásica')
bien('<conditionalFormatting ' in hoja, 'los colores quedan en la forma clásica')
bien('x14:dataValidation' not in hoja, 'la lista ya no está duplicada en la extensión')
bien('x14:conditionalFormatting' not in hoja, 'los colores ya no están duplicados')
bien('showDropDown="0"' in hoja, 'la flechita de la lista se ve fuera del Drive')
bien(hoja.find('<conditionalFormatting') < hoja.find('<dataValidations'),
     'el orden del esquema: colores antes que validaciones')
bien(hoja.find('<dataValidations') < hoja.find('<pageMargins'),
     'las validaciones van antes de pageMargins, como manda el esquema')

# ── después: openpyxl (léase Google o LibreOffice) sí las ve ─────────────
with warnings.catch_warnings():
    warnings.simplefilter('ignore')
    libro = load_workbook(SALIDA)
h = libro.worksheets[0]
validaciones = h.data_validations.dataValidation
colores = list(h.conditional_formatting)
print('después de la pasada: validaciones=%d colores=%d' % (len(validaciones), len(colores)))
bien(len(validaciones) >= 1, 'un programa de fuera ya ve la lista')
bien(len(colores) >= 1, 'un programa de fuera ya ve los colores')
if validaciones:
    dv = validaciones[0]
    bien(dv.type == 'list' and 'Opción' in (dv.formula1 or ''),
         'la lista trae sus valores: %s' % (dv.formula1 or '')[:40])
rellenos = []
for rango in colores:
    for regla in rango.rules:
        relleno = getattr(regla.dxf, 'fill', None) if regla.dxf else None
        if relleno is not None and relleno.bgColor is not None:
            rellenos.append(str(relleno.bgColor.rgb))
print('  colores encontrados:', rellenos)
bien(any('D32F2F' in c.upper() for c in rellenos), 'el rojo de «Opción 1» llega intacto')
bien(any('1565C0' in c.upper() for c in rellenos), 'el azul de «Opción 2» llega intacto')

# ── un archivo sin nada que traducir no se toca ──────────────────────────
bien(compat.convertir(AQUI + '/fixture-normal.xlsx', '/tmp/compat-nada.xlsx') is False,
     'un archivo normal no se reescribe (no hay nada que traducir)')

print()
print('✓ %d comprobaciones, %d en rojo' % (verde + rojo, rojo))
sys.exit(1 if rojo else 0)
