# -*- coding: utf-8 -*-
"""Audita las pestañas de provincia: ¿alguna mira datos que no son suyos?

Tres cruces posibles, las tres vistas de verdad el 22/09/2026 al clonar BOLIVAR:

  1. **Fórmula de otra tabla**: los TOTAL y la DURACIÓN nombran la tabla de su
     hoja (`Resp_BOLIVAR[[#This Row],[…]]`). Si una pestaña nombra la de otra,
     está sumando las filas de la otra.
  2. **Columna de territorio equivocada**: la columna «Territorio - X» debe
     traer la columna «Territorio - X» de «Prueba IFO», no la de otra provincia.
  3. **Celda de servicio**: `CB4` tiene que ser la provincia de la pestaña, tal
     como la escribe el formulario, o la pestaña traería las filas de otra.

Además comprueba que cada columna del formulario apunta a la columna de la hoja
base que dice su encabezado, y que no quedan celdas con error.

Uso:  python auditar_provincias.py <usuario> "<ruta>"

Autoría: Equipo de Tecnología Maquita — 2026-09-22
"""
import re
import sys

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
sys.path.insert(0, '/home/sistemas/almacen-maquita/herramientas')

import openpyxl

import nucleo_archivos as nucleo
from encuestas_hoja_xml import _col_letras, _col_num
from formulas_provincias import (HOJA_BASE, FILA_CABECERAS_BASE, FILA_CABECERAS_PESTANA,
                                 CLAVES, EQUIVALENCIAS, plano)

RE_INDICE = re.compile(r"INDEX\('?" + re.escape(HOJA_BASE) + r"'?!\$?([A-Z]{1,3})\$?\d")
RE_TABLA = re.compile(r"(Resp_[A-Za-z0-9_]+)\[")
PRIMERA = FILA_CABECERAS_PESTANA + 1


def auditar(fisica):
    libro = openpyxl.load_workbook(fisica)
    valores = openpyxl.load_workbook(fisica, data_only=True)
    base = libro[HOJA_BASE]
    titulo_base = {}
    por_nombre = {}
    for c in range(1, 120):
        valor = base.cell(row=FILA_CABECERAS_BASE, column=c).value
        if valor:
            titulo_base[c] = str(valor)
            por_nombre.setdefault(plano(valor), c)

    problemas = 0
    for nombre in libro.sheetnames:
        if nombre == HOJA_BASE:
            continue
        hoja, vista = libro[nombre], valores[nombre]
        avisos = []
        propia = (list(hoja.tables) or [None])[0]
        clave_esperada = {plano(k): v for k, v in CLAVES.items()}.get(plano(nombre))
        clave = hoja['CB4'].value
        if clave_esperada and plano(clave) != plano(clave_esperada):
            avisos.append('CB4 es «%s» y debería ser «%s»' % (clave, clave_esperada))

        ajenas, mal_apuntadas = {}, []
        for columna in range(1, 120):
            titulo = hoja.cell(row=FILA_CABECERAS_PESTANA, column=columna).value
            if not titulo:
                continue
            esperada = por_nombre.get(EQUIVALENCIAS.get(plano(titulo), plano(titulo)))
            for fila in range(PRIMERA, PRIMERA + 100):
                formula = hoja.cell(row=fila, column=columna).value
                if not (isinstance(formula, str) and formula.startswith('=')):
                    continue
                for tabla in RE_TABLA.findall(formula):
                    if propia and tabla != propia:
                        ajenas.setdefault(tabla, set()).add(_col_letras(columna))
                apunta = RE_INDICE.search(formula)
                if apunta and esperada and _col_num(apunta.group(1)) != esperada:
                    aviso = ('%s («%s») trae la columna %s de la base, que es «%s»'
                             % (_col_letras(columna), str(titulo)[:28], apunta.group(1),
                                titulo_base.get(_col_num(apunta.group(1)), '?')[:28]))
                    if aviso not in mal_apuntadas:
                        mal_apuntadas.append(aviso)
        for tabla, columnas in ajenas.items():
            avisos.append('fórmulas de la tabla %s en las columnas %s'
                          % (tabla, ', '.join(sorted(columnas))))
        avisos += mal_apuntadas[:6]
        errores = sum(1 for f in range(PRIMERA, PRIMERA + 100) for c in range(1, 60)
                      if isinstance(vista.cell(row=f, column=c).value, str)
                      and vista.cell(row=f, column=c).value.startswith('#'))
        if errores:
            avisos.append('%d celdas con error' % errores)
        if avisos:
            problemas += 1
            print('  %-12s PROBLEMAS' % nombre)
            for aviso in avisos:
                print('      - %s' % aviso)
        else:
            print('  %-12s bien (tabla %s, CB4 «%s»)' % (nombre, propia, clave))
    return problemas


if __name__ == '__main__':
    ruta = sys.argv[2]
    fallos = auditar(nucleo.ruta_fisica(int(sys.argv[1]), ruta))
    print('\n%s' % ('TODO CORRECTO' if not fallos else '%d pestañas con problemas' % fallos))
    sys.exit(1 if fallos else 0)
