# -*- coding: utf-8 -*-
"""Reparte en la hoja «Prueba IFO» las dos columnas agrupadas antiguas y las borra.

Hasta el 22/09/2026 la hoja de respuestas tenía dos columnas de cuando esas
preguntas eran una sola:

  · «¿Qué tipo de mecanismos utiliza la Organización comunitaria…» — casillas,
    con los enunciados marcados separados por comas.
  · «¿En qué ámbito?» — cuadrícula, «Fomento productivo: Sí · Ecologia…».

El formulario ya las tiene desglosadas en ocho preguntas de Sí/No, y la hoja ya
tiene una columna para cada una (vacías en las respuestas antiguas, porque el
dato estaba dentro de las agrupadas). Esto lo reparte:

  · casillas → «Sí» si su enunciado está marcado, «No» si no lo está;
  · cuadrícula → el valor que trae cada fila; lo que la persona no contestó se
    queda vacío (no responder no es «No»).

Y después BORRA las dos columnas agrupadas. Se hace con el editor (docbuilder),
no con openpyxl, para que las fórmulas de las pestañas de provincia se ajusten
solas al desplazarse las columnas.

Uso:
    python base_desglosar_agrupadas.py <usuario> "<ruta>" [--aplicar]

Autoría: Equipo de Tecnología Maquita — 2026-09-22
"""
import io
import os
import re
import secrets
import sys
import time
import unicodedata

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
sys.path.insert(0, '/home/sistemas/almacen-maquita/herramientas')

import openpyxl
import requests

import nucleo_archivos as nucleo
from api_onlyoffice import firmar_jwt, url_interna_ds, _reescribir_url_interna
from config_almacen import URL_PUBLICA
from encuestas_hoja_recalculo import CARPETA_SCRIPTS, URL_SCRIPTS
from formulas_provincias import HOJA_BASE, FILA_CABECERAS_BASE, letra, plano, _js

AGRUPADA_CASILLAS = '¿QUE TIPO DE MECANISMOS UTILIZA LA ORGANIZACION COMUNITARIA PARA ' \
                    'GARANTIZAR LA TRANSPARENCIA Y LA RENDICION DE CUENTAS?'
AGRUPADA_CUADRICULA = '¿EN QUE AMBITO?'

# Huella de cada enunciado: un trozo que aparece igual en el dato histórico
# aunque la primera palabra cambie («Pose» en el formulario, «Posee» en la
# respuesta que se guardó entonces).
HUELLAS = {
    'REALIZA AL MENOS 1 REUNION/ASAMBLEA': 'AL MENOS 1 REUNION/ASAMBLEA',
    'POSE UN INVENTARIO ACTUALIZADO': 'UN INVENTARIO ACTUALIZADO DE BIENES',
    'CUENTA CON DOCUMENTOS DE RESPALDO': 'CON DOCUMENTOS DE RESPALDO',
}


def huella_de(titulo):
    normal = plano(titulo)
    for prefijo, huella in HUELLAS.items():
        if normal.startswith(prefijo):
            return huella
    return normal[:30]


def reparto_casillas(valor, titulos):
    """{título: 'Sí'/'No'} según los enunciados marcados. Celda vacía → nada."""
    if valor in (None, ''):
        return {}
    texto = plano(valor)
    return {t: ('Sí' if huella_de(t) in texto else 'No') for t in titulos}


def reparto_cuadricula(valor, titulos):
    """{título: valor} leyendo «Fila: valor · Fila: valor». Lo no contestado se
    queda fuera: una fila sin responder no es un «No»."""
    if valor in (None, ''):
        return {}
    por_nombre = {}
    for parte in str(valor).split(' · '):
        nombre, separador, dato = parte.rpartition(': ')
        if separador:
            por_nombre[plano(nombre)] = dato.strip()
    return {t: por_nombre[plano(t)] for t in titulos if plano(t) in por_nombre}


def planear(fisica):
    """(instrucciones de escritura, columnas a borrar, cuántas celdas)."""
    libro = openpyxl.load_workbook(fisica, data_only=True)
    hoja = libro[HOJA_BASE]
    titulos, por_nombre = {}, {}
    for c in range(1, 120):
        valor = hoja.cell(row=FILA_CABECERAS_BASE, column=c).value
        if valor:
            titulos[c] = str(valor)
            por_nombre.setdefault(plano(valor), c)

    col_casillas = por_nombre.get(AGRUPADA_CASILLAS)
    col_cuadricula = por_nombre.get(AGRUPADA_CUADRICULA)
    if not col_casillas and not col_cuadricula:
        return [], [], 0

    # Las columnas de destino son las que ya creó el motor, por su nombre.
    destino_casillas = [t for t in titulos.values() if huella_de(t) in HUELLAS.values()]
    destino_cuadricula = ['Fomento productivo', 'Ecologia - Ambiente', 'Comercialización',
                          'Derechos y Género', 'Servicios básicos']
    destino_cuadricula = [t for t in destino_cuadricula if plano(t) in por_nombre]

    escrituras, celdas = [], 0
    for fila in range(FILA_CABECERAS_BASE + 1, hoja.max_row + 1):
        if hoja.cell(row=fila, column=1).value in (None, ''):
            continue
        reparto = {}
        if col_casillas:
            reparto.update(reparto_casillas(
                hoja.cell(row=fila, column=col_casillas).value, destino_casillas))
        if col_cuadricula:
            reparto.update(reparto_cuadricula(
                hoja.cell(row=fila, column=col_cuadricula).value, destino_cuadricula))
        for titulo, dato in reparto.items():
            columna = por_nombre[plano(titulo)]
            if hoja.cell(row=fila, column=columna).value in (None, ''):
                escrituras.append((columna, fila, dato))
                celdas += 1
    borrar = sorted([c for c in (col_casillas, col_cuadricula) if c], reverse=True)
    return escrituras, borrar, celdas


def principal(usuario, ruta, aplicar):
    fisica = nucleo.ruta_fisica(usuario, ruta)
    escrituras, borrar, celdas = planear(fisica)
    print('  %d celdas que repartir, columnas que borrar: %s' % (celdas, borrar))
    if not escrituras and not borrar:
        print('  nada que hacer')
        return 0

    token = firmar_jwt({'u': usuario, 'r': ruta, 'uso': 'descarga',
                        'exp': int(time.time()) + 900})
    url_libro = '%s/api/almacen/onlyoffice/download?t=%s' % (URL_PUBLICA, token)
    lineas = ['builder.OpenFile("%s");' % url_libro,
              'var hoja = Api.GetSheet(%s);' % _js(HOJA_BASE), 'if (hoja) {']
    for columna, fila, dato in escrituras:
        lineas.append('  hoja.GetRange("%s%d").SetValue(%s);'
                      % (letra(columna), fila, _js(dato)))
    # De derecha a izquierda: borrar una columna no descuadra la siguiente.
    for columna in borrar:
        lineas.append('  hoja.GetRange("%s:%s").Delete();' % (letra(columna), letra(columna)))
    lineas += ['}', 'Api.RecalculateAllFormulas();',
               'builder.SaveFile("xlsx", "base.xlsx");', 'builder.CloseFile();']

    os.makedirs(CARPETA_SCRIPTS, exist_ok=True)
    fichero = secrets.token_hex(16) + '.docbuilder'
    camino = os.path.join(CARPETA_SCRIPTS, fichero)
    open(camino, 'w', encoding='utf-8').write('\n'.join(lineas) + '\n')
    os.chmod(camino, 0o644)
    print('  guion de %d líneas' % len(lineas))
    try:
        cuerpo = {'async': False, 'key': 'base' + secrets.token_hex(8),
                  'url': '%s%s/%s' % (URL_PUBLICA, URL_SCRIPTS, fichero)}
        cuerpo['token'] = firmar_jwt(dict(cuerpo))
        respuesta = requests.post(
            url_interna_ds().rstrip('/') + '/docbuilder', json=cuerpo, timeout=(3, 900),
            headers={'Authorization': 'Bearer ' + firmar_jwt({'payload': cuerpo})})
        try:
            urls = list((respuesta.json() or {}).get('urls', {}).values())
        except ValueError:
            urls = []
        if not urls:
            print('  el docbuilder respondió:', respuesta.status_code, repr(respuesta.text[:500]))
            return 1
        salida = requests.get(_reescribir_url_interna(urls[0]), timeout=(3, 300)).content
        destino = '/tmp/base-desglosada.xlsx'
        open(destino, 'wb').write(salida)
        print('  resultado en', destino, '(%.0f KB)' % (len(salida) / 1024))
        if aplicar:
            import sala_editor
            from api_onlyoffice import _base_documento, invalidar_cache
            dentro = sala_editor.usuarios_conectados(_base_documento(usuario, ruta))
            if dentro is None or dentro:
                print('  NO se aplica: el Excel está abierto en el editor (%s)' % dentro)
                return 2
            carpeta, _, nombre = ruta.rpartition('/')
            nucleo.subir(usuario, carpeta or '/', nombre, io.BytesIO(salida))
            invalidar_cache(usuario, ruta)
            print('  aplicado al Drive (la versión anterior queda en el historial)')
    finally:
        os.remove(camino)
    return 0


if __name__ == '__main__':
    sys.exit(principal(int(sys.argv[1]), sys.argv[2], '--aplicar' in sys.argv))
