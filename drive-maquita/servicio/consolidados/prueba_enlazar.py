#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prueba de `formularios_libro.enlazar`, en la carpeta de pruebas del usuario 14.

Comprueba lo que pidió Wilson:
  · al enlazar un formulario a un libro se crea la hoja «Respuestas 1» y se llena;
  · un SEGUNDO formulario sobre el mismo libro crea «Respuestas 2», sin tocar la
    primera;
  · las fórmulas propias del libro conservan su resultado.
"""
import io
import sys
import uuid

import openpyxl

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')

import almacen_bd as bd  # noqa: E402
import encuestas_bd as ebd  # noqa: E402
import encuestas_hoja as hoja_mod  # noqa: E402
import formularios_libro as fl  # noqa: E402
import nucleo_archivos as nucleo  # noqa: E402

USUARIO = 14
CARPETA = '/PRUEBAS FORMULARIOS'
LIBRO = CARPETA + '/Libro de trabajo demo.xlsx'
FISICO = '/mnt/almacen/%d/archivos' % USUARIO


def _crear_respuestas(nombre, cabeceras, filas):
    ruta = CARPETA + '/' + nombre
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Respuestas'
    hoja.append(cabeceras)
    for fila in filas:
        hoja.append(fila)
    memoria = io.BytesIO()
    libro.save(memoria)
    libro.close()
    memoria.seek(0)
    nucleo.subir(USUARIO, CARPETA, nombre, memoria)
    return ruta


def _registrar_formulario(titulo, ruta_respuestas):
    encuesta_id = uuid.uuid4().hex
    ebd.registrar(encuesta_id, USUARIO, CARPETA + '/' + titulo + '.forma', titulo)
    hoja_mod.vincular(encuesta_id, ruta_respuestas)
    return encuesta_id


def main():
    # Se parte de un libro limpio, con una fórmula propia.
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Mis datos'
    hoja['A1'] = 10
    hoja['A2'] = 32
    hoja['A3'] = '=A1+A2'
    memoria = io.BytesIO()
    libro.save(memoria)
    libro.close()
    memoria.seek(0)
    nucleo.subir(USUARIO, CARPETA, 'Libro de trabajo demo.xlsx', memoria)
    bd.ejecutar('DELETE FROM vinculos_datos WHERE destino_ruta = %s', (LIBRO,))
    print('libro de trabajo limpio, con la fórmula A3 = A1+A2')

    r1 = _crear_respuestas('Encuesta clima (respuestas).xlsx',
                           ['Fecha', 'Nombre', 'Le gusta'],
                           [['2026-09-09 16:00', 'Ana', 'Sí']])
    id1 = _registrar_formulario('Encuesta clima', r1)
    hoja1 = fl.enlazar(USUARIO, id1, LIBRO)
    print('formulario 1 enlazado a la hoja: %s' % hoja1)

    r2 = _crear_respuestas('Encuesta transporte (respuestas).xlsx',
                           ['Fecha', 'Nombre', 'Medio', 'Minutos'],
                           [['2026-09-09 16:05', 'Beto', 'Bus', 45]])
    id2 = _registrar_formulario('Encuesta transporte', r2)
    hoja2 = fl.enlazar(USUARIO, id2, LIBRO)
    print('formulario 2 enlazado a la hoja: %s' % hoja2)

    print('\n--- cómo quedó el libro ---')
    wb = openpyxl.load_workbook(FISICO + LIBRO, data_only=True)
    print('hojas:', wb.sheetnames)
    for nombre in (hoja1, hoja2):
        print('  [%s]' % nombre)
        for fila in wb[nombre].iter_rows(min_row=1, max_row=3, max_col=4,
                                         values_only=True):
            if any(v is not None for v in fila):
                print('     ', fila)
    print('  fórmula propia A3 (debe valer 42):', wb['Mis datos']['A3'].value)
    wb.close()

    print('\n--- hojas del libro que reciben respuestas ---')
    for enlace in fl.hojas_enlazadas(LIBRO):
        print('  %s  ←  %s' % (enlace['destino_hoja'],
                               enlace['origen_ruta'].split('/')[-1]))


if __name__ == '__main__':
    main()
