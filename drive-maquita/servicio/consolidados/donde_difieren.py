#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""¿Las diferencias de los consolidados están SOLO en el bloque consolidado?

Los cuatro consolidados difieren del archivo de Google porque los regeneramos a
propósito. La pregunta que importa es otra: ¿esa reescritura tocó algo fuera del
bloque —la tabla de configuración, las otras hojas, lo que la gente escribió a
mano—? Si toda la diferencia cae dentro del bloque, no se perdió trabajo de
nadie.
"""
import sys

import openpyxl

BLOQUES = {
    'Mapa general de inversiones 2026.xlsx': ('ConsolidadoP', 4, 8),
    'Mapa general de inversiones 2025.xlsx': ('ConsolidadoP', 4, 8),
    'Respaldo mapa inversiones.xlsx': ('ConsolidadoP', 4, 8),
    'Mapa general Ejec. Técnica 2026.xlsx': ('ConsolidadoT', 4, 7),
}

CRUDO = '/mnt/almacen/_sincronizacion-asc/crudo/mi-unidad/Mapa de inversiones/'
PUB = ('/mnt/almacen/_unidades/12/archivos/Mi unidad de Google/'
       'Mapa de inversiones/')


def valores(ruta):
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    datos = {}
    try:
        for nombre in libro.sheetnames:
            for fila in libro[nombre].iter_rows():
                for celda in fila:
                    if celda.value is not None:
                        datos[(nombre, celda.row, celda.column)] = celda.value
    finally:
        libro.close()
    return datos


def main():
    for nombre, (hoja_bloque, fila0, col0) in BLOQUES.items():
        try:
            antes = valores(CRUDO + nombre)
            despues = valores(PUB + nombre)
        except Exception as excepcion:
            print('%-42s ERROR: %s' % (nombre[:42], excepcion))
            continue

        difieren = [k for k in antes
                    if k not in despues or antes[k] != despues[k]]
        dentro = [k for k in difieren
                  if k[0] == hoja_bloque and k[1] >= fila0 and k[2] >= col0]
        fuera = [k for k in difieren if k not in set(dentro)]

        print('== %s' % nombre)
        print('   celdas que difieren      : %d' % len(difieren))
        print('   DENTRO del bloque        : %d' % len(dentro))
        print('   FUERA del bloque (grave) : %d' % len(fuera))
        for clave in fuera[:6]:
            print('       hoja %r fila %s col %s: Google=%r Maquita=%r'
                  % (clave[0], clave[1], clave[2],
                     antes.get(clave), despues.get(clave)))
        print()


if __name__ == '__main__':
    main()
