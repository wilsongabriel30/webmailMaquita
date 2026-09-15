#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prueba del circuito «respuestas → hoja del libro», en la carpeta de pruebas.

Valida, con archivos reales pero en `/PRUEBAS FORMULARIOS` del usuario 14:

  1. que un vínculo lleva los datos del archivo de respuestas a una hoja del
     libro de trabajo;
  2. que al llegar una respuesta nueva la hoja se actualiza SOLA, gracias al
     aviso que se añadió hoy en `encuestas_hoja`;
  3. si al refrescar el vínculo el libro de trabajo PIERDE los resultados de sus
     propias fórmulas — el mismo defecto que apareció en los consolidados.

Uso:  prueba_formulario_hoja.py preparar|responder|ver
"""
import io
import sys

import openpyxl

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')

import almacen_bd as bd  # noqa: E402
import nucleo_archivos as nucleo  # noqa: E402
import api_vinculos as vinc  # noqa: E402

USUARIO = 14
CARPETA = '/PRUEBAS FORMULARIOS'
ORIGEN = CARPETA + '/Respuestas demo.xlsx'
DESTINO = CARPETA + '/Libro de trabajo demo.xlsx'
HOJA_ORIGEN = 'Respuestas'
HOJA_DESTINO = 'Respuestas 1'
FISICO = '/mnt/almacen/%d/archivos' % USUARIO


def _subir(ruta, libro):
    memoria = io.BytesIO()
    libro.save(memoria)
    libro.close()
    memoria.seek(0)
    carpeta, _, nombre = ruta.rpartition('/')
    nucleo.subir(USUARIO, carpeta or '/', nombre, memoria)


def preparar():
    # Archivo de respuestas: lo que hoy genera el formulario.
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = HOJA_ORIGEN
    hoja.append(['Fecha', 'Nombre', 'Provincia', 'Cuántos'])
    hoja.append(['2026-09-09 10:00', 'Ana', 'Manabí', 3])
    _subir(ORIGEN, libro)
    print('creado: %s' % ORIGEN)

    # Libro de trabajo de la persona: con datos y una FÓRMULA propia, para ver
    # si el refresco del vínculo se la deja sin resultado.
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Mis datos'
    hoja['A1'] = 10
    hoja['A2'] = 32
    hoja['A3'] = '=A1+A2'          # debe seguir valiendo 42
    libro.create_sheet(HOJA_DESTINO)
    _subir(DESTINO, libro)
    print('creado: %s (con la fórmula A3 = A1+A2)' % DESTINO)

    bd.ejecutar('DELETE FROM vinculos_datos WHERE destino_ruta = %s', (DESTINO,))
    fila = bd.ejecutar(
        """INSERT INTO vinculos_datos
           (origen_usuario, origen_ruta, origen_hoja, origen_rango,
            destino_usuario, destino_ruta, destino_hoja, destino_celda, creado_por)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING *""",
        (USUARIO, ORIGEN, HOJA_ORIGEN, 'A1:D200',
         USUARIO, DESTINO, HOJA_DESTINO, 'A1', USUARIO))
    ok, mensaje = vinc._refrescar(dict(fila))
    print('vínculo creado y materializado: %s %s' % (ok, mensaje))


def responder():
    """Simula que entra una respuesta nueva, igual que hace el formulario."""
    libro = openpyxl.load_workbook(FISICO + ORIGEN)
    libro[HOJA_ORIGEN].append(['2026-09-09 15:40', 'Beto', 'Esmeraldas', 7])
    _subir(ORIGEN, libro)
    print('respuesta añadida al archivo de respuestas')
    # Esto es lo que ahora hace encuestas_hoja al entrar una respuesta.
    vinc.refrescar_por_origen(USUARIO, ORIGEN)
    print('vínculos avisados')


def ver():
    libro = openpyxl.load_workbook(FISICO + DESTINO, data_only=True)
    hoja = libro[HOJA_DESTINO]
    print('--- hoja «%s» del libro de trabajo ---' % HOJA_DESTINO)
    for fila in hoja.iter_rows(min_row=1, max_row=4, max_col=4, values_only=True):
        if any(v is not None for v in fila):
            print('   ', fila)
    mis = libro['Mis datos']
    print('--- la fórmula propia del libro ---')
    print('    A3 (debe valer 42):', repr(mis['A3'].value))
    libro.close()
    libro = openpyxl.load_workbook(FISICO + DESTINO)
    print('    A3 sigue siendo fórmula:', repr(libro['Mis datos']['A3'].value))
    libro.close()


if __name__ == '__main__':
    accion = sys.argv[1] if len(sys.argv) > 1 else 'ver'
    {'preparar': preparar, 'responder': responder, 'ver': ver}[accion]()
