#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Devuelve a las fórmulas su resultado guardado, sin tocar nada más.

El problema. `openpyxl` conserva las fórmulas pero descarta su resultado
guardado, así que los totales de cabecera de los consolidados quedaban en
blanco. Un `.xlsx` guarda cada celda con fórmula así:

    <c r="J1"><f>V3+AB3+AH3</f><v>5144476.638</v></c>
                                ^^^^^^^^^^^^^^^^^^  esto es lo que faltaba

Los dos intentos anteriores, descartados:
  · LibreOffice recalcula bien, pero al reescribir el libro convierte el texto
    `#REF!` en la fórmula `=#ref!` y pierde anchos de columna.
  · El Document Server de OnlyOffice no estropea nada, pero tampoco recalcula:
    devuelve el archivo igual de vacío.

Lo que hace este módulo. Usa LibreOffice SOLO como calculadora, sobre una copia
desechable, y del resultado se queda únicamente con los números. Después los
inyecta en el XML del archivo bueno, celda por celda, añadiendo el `<v>` que
falta. El archivo publicado sigue siendo el nuestro: mismas fórmulas, mismo
formato, mismos anchos. Solo se le añade lo que faltaba.
"""
import os
import re
import shutil
import sys
import tempfile
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from recalcular import recalcular  # noqa: E402  (LibreOffice, como calculadora)

# <c r="J1" s="12"> … </c>  o  <c r="J1" s="12"/>
_RE_CELDA = re.compile(rb'<c r="([A-Z]+\d+)"([^>]*?)(/>|>(.*?)</c>)', re.S)


def _valores_calculados(ruta):
    """{(hoja, referencia): valor} de las celdas que tienen fórmula CON resultado."""
    libro_v = openpyxl.load_workbook(ruta, data_only=True)
    libro_f = openpyxl.load_workbook(ruta, data_only=False)
    calculados = {}
    try:
        for hoja in libro_f.sheetnames:
            pf, pv = libro_f[hoja], libro_v[hoja]
            for fila in pf.iter_rows():
                for celda in fila:
                    if isinstance(celda.value, str) and celda.value.startswith('='):
                        valor = pv.cell(row=celda.row, column=celda.column).value
                        if valor is not None:
                            referencia = '%s%d' % (get_column_letter(celda.column),
                                                   celda.row)
                            calculados[(hoja, referencia)] = valor
        return calculados
    finally:
        libro_f.close()
        libro_v.close()


def _hojas_del_paquete(ruta):
    """nombre de hoja -> ruta interna del xml, respetando el orden del libro."""
    libro = openpyxl.load_workbook(ruta, read_only=True)
    nombres = list(libro.sheetnames)
    libro.close()
    with zipfile.ZipFile(ruta) as paquete:
        internas = sorted(
            (n for n in paquete.namelist()
             if re.match(r'xl/worksheets/sheet\d+\.xml$', n)),
            key=lambda n: int(re.search(r'(\d+)\.xml$', n).group(1)))
    return dict(zip(nombres, internas))


def _texto_valor(valor):
    """Representación del valor y el tipo que le corresponde en el XML."""
    if isinstance(valor, bool):
        return b'1' if valor else b'0', b' t="b"'
    if isinstance(valor, (int, float)):
        texto = repr(float(valor)) if isinstance(valor, float) else str(valor)
        return texto.encode(), b''
    texto = str(valor)
    if texto.startswith('#'):
        # NUNCA se inyecta un error. `#NAME?` suele significar solo que la
        # calculadora no conoce una función que el editor sí entiende: escribirlo
        # convertiría una celda vacía en una celda con error, que es peor de lo
        # que había. Si la fórmula da error de verdad, el editor lo mostrará al
        # abrir; no hace falta grabarlo.
        return None, None
    return None, None                          # texto: se deja como está


def inyectar(ruta_destino, calculados, hojas):
    """Añade <v> a las celdas con fórmula que no lo tienen. Devuelve cuántas."""
    puestos = 0
    temporal = ruta_destino + '.tmp'
    with zipfile.ZipFile(ruta_destino) as entrada:
        with zipfile.ZipFile(temporal, 'w', zipfile.ZIP_DEFLATED,
                             compresslevel=6) as salida:
            internas = {v: k for k, v in hojas.items()}
            for elemento in entrada.infolist():
                datos = entrada.read(elemento.filename)
                hoja = internas.get(elemento.filename)
                if hoja:
                    datos, puestas = _inyectar_en_hoja(datos, hoja, calculados)
                    puestos += puestas
                salida.writestr(elemento, datos)
    os.replace(temporal, ruta_destino)
    return puestos


def _inyectar_en_hoja(datos, hoja, calculados):
    puestas = 0

    def reemplazo(coincidencia):
        nonlocal puestas
        referencia = coincidencia.group(1).decode()
        atributos = coincidencia.group(2)
        cuerpo = coincidencia.group(4)
        # Solo celdas con fórmula a las que les falta el resultado. Ojo: openpyxl
        # no omite el elemento, lo deja VACÍO (`<v></v>`), así que no basta con
        # mirar si existe `<v>`; hay que mirar si trae algo dentro.
        if cuerpo is None or b'<f' not in cuerpo:
            return coincidencia.group(0)
        if b'<v>' in cuerpo and b'<v></v>' not in cuerpo:
            return coincidencia.group(0)
        valor = calculados.get((hoja, referencia))
        if valor is None:
            return coincidencia.group(0)
        texto, tipo = _texto_valor(valor)
        if texto is None:
            return coincidencia.group(0)
        if tipo and b' t="' not in atributos:
            atributos = atributos + tipo
        nuevo_cuerpo = (cuerpo.replace(b'<v></v>', b'<v>' + texto + b'</v>')
                        if b'<v></v>' in cuerpo
                        else cuerpo + b'<v>' + texto + b'</v>')
        puestas += 1
        return b'<c r="' + referencia.encode() + b'"' + atributos + b'>' \
               + nuevo_cuerpo + b'</c>'

    return _RE_CELDA.sub(reemplazo, datos), puestas


def completar(ruta_archivo):
    """Rellena los resultados que falten en `ruta_archivo`. (ok, puestos, motivo)."""
    carpeta = tempfile.mkdtemp(prefix='calcular-')
    try:
        copia = os.path.join(carpeta, 'copia.xlsx')
        calculada = os.path.join(carpeta, 'calculada.xlsx')
        shutil.copy2(ruta_archivo, copia)

        ok, motivo = recalcular(copia, calculada)
        if not ok:
            return False, 0, motivo

        calculados = _valores_calculados(calculada)
        if not calculados:
            return False, 0, 'la calculadora no devolvió ningún resultado'

        hojas = _hojas_del_paquete(ruta_archivo)
        puestos = inyectar(ruta_archivo, calculados, hojas)
        return True, puestos, ''
    except Exception as excepcion:
        return False, 0, str(excepcion)[:200]
    finally:
        shutil.rmtree(carpeta, ignore_errors=True)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit('uso: inyectar_valores.py <archivo.xlsx>')
    ok, puestos, motivo = completar(sys.argv[1])
    print('OK: %d resultados añadidos' % puestos if ok else 'FALLO: %s' % motivo)
    sys.exit(0 if ok else 1)
