# -*- coding: utf-8 -*-
"""Deja las respuestas de un formulario en una hoja del libro de trabajo.

Es lo que hace Google Sheets: desde un Excel se crea un formulario y las
respuestas van cayendo en una hoja nueva de ese mismo libro; si se crea un
segundo formulario, se crea otra hoja.

Cómo se consigue aquí, sin reescribir el libro en cada respuesta —que pisaría a
quien lo tenga abierto—:

  1. las respuestas siguen acumulándose en el archivo que ya genera el
     formulario (`<título> (respuestas).xlsx`), como hasta ahora;
  2. en el libro de trabajo se crea una hoja con EL NOMBRE DEL FORMULARIO
     (hasta el 10/09/2026 se llamaba «Respuestas N»; ver `formularios_nombres`);
  3. un vínculo de datos trae ese archivo a esa hoja, y `encuestas_hoja` avisa a
     los vínculos con cada respuesta, así que la hoja se llena sola.

Para la persona el efecto es el de Google. Por debajo, su libro solo se escribe
cuando llega una respuesta, nunca a mitad de una edición suya.
"""
import logging

import openpyxl

import almacen_bd as bd
import encuestas_bd as ebd
import encuestas_hoja as hoja_mod
import formularios_nombres as nombres
from seguridad_rutas import ruta_fisica

log = logging.getLogger(__name__)

PREFIJO_HOJA = 'Respuestas'
# Sitio reservado en la hoja del libro. Si un formulario supera estas filas, las
# respuestas siguen guardadas y en su archivo: lo que se queda corto es la vista.
FILAS_RESERVADAS = 5000
# Columnas reservadas. NO se usa el ancho que tiene la hoja de respuestas al
# crearla: el formulario nace casi vacío y la gente le añade preguntas después.
# Si el rango se fijara al ancho inicial, las preguntas nuevas no se verían en la
# hoja del libro. Se reserva sitio de sobra y las columnas sin usar quedan vacías.
COLUMNAS_RESERVADAS = 80


class SinHoja(Exception):
    """El formulario todavía no tiene archivo de respuestas."""


def hojas_del_libro(usuario, ruta_libro):
    """Nombres de las hojas que tiene hoy el libro."""
    fisica = ruta_fisica(usuario, ruta_libro)
    libro = openpyxl.load_workbook(fisica, read_only=True)
    try:
        return set(libro.sheetnames)
    finally:
        libro.close()


def hoja_libre(usuario, ruta_libro, titulo=None):
    """Nombre de hoja para este formulario que no exista ya en el libro.

    Con título: «Encuesta de clima», «Encuesta de clima (2)»… (como Google).
    Sin título (llamadas antiguas): «Respuestas N».
    """
    usadas = hojas_del_libro(usuario, ruta_libro)
    if titulo:
        return nombres.hoja_libre(usadas, nombres.nombre_hoja(titulo))
    numero = 1
    while '%s %d' % (PREFIJO_HOJA, numero) in usadas:
        numero += 1
    return '%s %d' % (PREFIJO_HOJA, numero)


def _ancho_respuestas(usuario, ruta_respuestas):
    """(columnas, nombre de la hoja, fila de cabeceras) de la tabla de respuestas.

    El archivo maestro lleva el logo institucional arriba y la tabla debajo
    (`encuestas_excel.escribir`): las cabeceras no están en la fila 1. La hoja
    del libro empieza EN las cabeceras (10/09/2026), como en Google, y así
    coinciden con las que el complemento escribe en la fila 1 del editor.
    """
    fisica = ruta_fisica(usuario, ruta_respuestas)
    libro = openpyxl.load_workbook(fisica, read_only=True)
    try:
        pagina = libro[libro.sheetnames[0]]
        fila_cabeceras = 1
        for fila in pagina.iter_rows(min_row=1, max_row=12, max_col=1,
                                     values_only=True):
            if fila and fila[0] not in (None, ''):
                break
            fila_cabeceras += 1
        else:
            fila_cabeceras = 1
        return max(pagina.max_column or 1, 1), libro.sheetnames[0], fila_cabeceras
    finally:
        libro.close()


def _crear_hoja_en_libro(usuario, ruta_libro, nombre_hoja):
    """Añade la hoja vacía al libro, si no está ya."""
    import io
    import nucleo_archivos as nucleo

    fisica = ruta_fisica(usuario, ruta_libro)
    libro = openpyxl.load_workbook(fisica)
    if nombre_hoja in libro.sheetnames:
        libro.close()
        return False
    libro.create_sheet(nombre_hoja)
    memoria = io.BytesIO()
    libro.save(memoria)
    libro.close()
    memoria.seek(0)
    carpeta, _, nombre = ruta_libro.rpartition('/')
    nucleo.subir(usuario, carpeta or '/', nombre, memoria)
    return True


def enlazar(usuario, encuesta_id, ruta_libro, nombre_hoja=None):
    """Crea la hoja en el libro y la deja vinculada a las respuestas.

    Devuelve el nombre de la hoja creada. Lanza SinHoja si el formulario aún no
    tiene archivo de respuestas (hay que exportarlo antes, una sola vez).
    """
    fila = ebd.obtener(encuesta_id)
    if not fila:
        raise SinHoja('No se encuentra el formulario')
    ruta_respuestas = hoja_mod.ruta_de(fila)
    if not ruta_respuestas:
        raise SinHoja('El formulario todavía no tiene archivo de respuestas')

    _, hoja_origen, fila_cabeceras = _ancho_respuestas(usuario, ruta_respuestas)
    nombre_hoja = nombre_hoja or hoja_libre(usuario, ruta_libro,
                                            fila.get('titulo'))
    _crear_hoja_en_libro(usuario, ruta_libro, nombre_hoja)

    from openpyxl.utils import get_column_letter
    rango = 'A%d:%s%d' % (fila_cabeceras, get_column_letter(COLUMNAS_RESERVADAS),
                          fila_cabeceras + FILAS_RESERVADAS - 1)

    # Un formulario, una hoja: si ya había vínculo para este destino, se rehace.
    bd.ejecutar("""DELETE FROM vinculos_datos
                    WHERE destino_ruta = %s AND destino_hoja = %s""",
                (ruta_libro, nombre_hoja))
    creado = bd.ejecutar(
        """INSERT INTO vinculos_datos
           (origen_usuario, origen_ruta, origen_hoja, origen_rango,
            destino_usuario, destino_ruta, destino_hoja, destino_celda, creado_por)
           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING *""",
        (int(fila['propietario']), ruta_respuestas, hoja_origen, rango,
         usuario, ruta_libro, nombre_hoja, 'A1', usuario))

    import api_vinculos
    ok, mensaje = api_vinculos._refrescar(dict(creado))
    if not ok:
        log.warning('formulario %s: hoja creada pero sin datos aún (%s)',
                    encuesta_id, mensaje)
    log.info('formulario %s enlazado a %s!%s', encuesta_id, ruta_libro, nombre_hoja)
    return nombre_hoja


def libro_de_respuestas(ruta_respuestas):
    """El libro al que alimenta este archivo de respuestas, o ''."""
    filas = bd.consultar(
        """SELECT destino_ruta FROM vinculos_datos
            WHERE origen_ruta = %s AND activo ORDER BY id DESC LIMIT 1""",
        (ruta_respuestas,))
    return filas[0]['destino_ruta'] if filas else ''


def hoja_de_respuestas(ruta_respuestas):
    """La hoja del libro que recibe este archivo de respuestas, o ''."""
    filas = bd.consultar(
        """SELECT destino_hoja FROM vinculos_datos
            WHERE origen_ruta = %s AND activo ORDER BY id DESC LIMIT 1""",
        (ruta_respuestas,))
    return filas[0]['destino_hoja'] if filas else ''


def hojas_enlazadas(ruta_libro):
    """Qué hojas de este libro reciben respuestas, y de qué formulario."""
    filas = bd.consultar(
        """SELECT destino_hoja, origen_ruta
             FROM vinculos_datos
            WHERE destino_ruta = %s AND activo
            ORDER BY destino_hoja""", (ruta_libro,))
    return [dict(f) for f in filas]
