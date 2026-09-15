#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prueba real: cambiar un dato en una matriz y ver si llega al consolidado.

Reproduce exactamente lo que ocurre cuando una persona edita desde el editor:
guarda con `nucleo.subir()` (lo mismo que hace el callback de OnlyOffice) y
después dispara el enganche. Luego comprueba el consolidado.

El valor original se restaura al final, pasando por el mismo camino.

Uso:  prueba_extremo_a_extremo.py marcar|restaurar|ver
"""
import io
import os
import sys
import time

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio/consolidados')

import openpyxl  # noqa: E402
import nucleo_archivos as nucleo  # noqa: E402

USUARIO = 14
BASE = '/mnt/almacen/_unidades/12/archivos/Mi unidad de Google'
MATRIZ_REL = '⚠️ Matrices de seguimiento/1 Carpeta Planificación Esmeraldas/Matriz Seguimiento Esmeraldas.xlsx'
MATRIZ_VIRTUAL = '/unidades/12/Mi unidad de Google/' + MATRIZ_REL
HOJA_ORIGEN = 'E7'
CELDA_ORIGEN = 'I4'          # primera fila del rango A4:I178

CONSOLIDADO = os.path.join(BASE, 'Mapa de inversiones/Mapa general de inversiones 2026.xlsx')
HOJA_DESTINO = 'ConsolidadoP'
CELDA_DESTINO = 'P4'         # H4 es la columna A del origen; I es la novena → P

RESGUARDO = '/mnt/almacen/_sincronizacion-asc/estado-consolidados/valor-original-prueba.txt'
MARCA = 'PRUEBA-CONSOLIDADO'


def escribir_en_matriz(valor):
    """Guarda el valor en la matriz igual que lo haría el editor."""
    ruta = os.path.join(BASE, MATRIZ_REL)
    libro = openpyxl.load_workbook(ruta)
    libro[HOJA_ORIGEN][CELDA_ORIGEN] = valor
    memoria = io.BytesIO()
    libro.save(memoria)
    libro.close()
    memoria.seek(0)
    carpeta, _, nombre = MATRIZ_VIRTUAL.rpartition('/')
    nucleo.subir(USUARIO, carpeta, nombre, memoria)


def leer_consolidado():
    libro = openpyxl.load_workbook(CONSOLIDADO, data_only=True)
    valor = libro[HOJA_DESTINO][CELDA_DESTINO].value
    libro.close()
    return valor


def leer_matriz():
    libro = openpyxl.load_workbook(os.path.join(BASE, MATRIZ_REL), data_only=True)
    valor = libro[HOJA_ORIGEN][CELDA_ORIGEN].value
    libro.close()
    return valor


def main():
    accion = sys.argv[1] if len(sys.argv) > 1 else 'ver'

    if accion == 'ver':
        print('matriz  %s!%s = %r' % (HOJA_ORIGEN, CELDA_ORIGEN, leer_matriz()))
        print('consolidado %s!%s = %r' % (HOJA_DESTINO, CELDA_DESTINO, leer_consolidado()))
        return

    if accion == 'marcar':
        original = leer_matriz()
        with open(RESGUARDO, 'w', encoding='utf-8') as fichero:
            fichero.write('' if original is None else str(original))
        nuevo = '%s %s' % (MARCA, time.strftime('%H:%M:%S'))
        print('valor original guardado: %r' % original)
        print('escribiendo en la matriz: %r' % nuevo)
        escribir_en_matriz(nuevo)
        from enganche import al_guardar
        print('enganche lanzado:', al_guardar(MATRIZ_VIRTUAL))
        return

    if accion == 'restaurar':
        try:
            with open(RESGUARDO, encoding='utf-8') as fichero:
                original = fichero.read()
        except OSError:
            sys.exit('no hay valor original guardado')
        original = original or None
        print('restaurando en la matriz: %r' % original)
        escribir_en_matriz(original)
        from enganche import al_guardar
        print('enganche lanzado:', al_guardar(MATRIZ_VIRTUAL))
        return

    sys.exit('uso: prueba_extremo_a_extremo.py marcar|restaurar|ver')


if __name__ == '__main__':
    main()
