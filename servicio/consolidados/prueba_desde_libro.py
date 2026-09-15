#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prueba del flujo «Crear formulario» tal como lo ejecuta el botón del editor.

Recorre los mismos pasos que el endpoint `formularios/desde-libro`, sin pasar por
Flask: crear el `.forma`, registrarlo, prepararle su archivo de respuestas y
dejar la hoja «Respuestas N» dentro del libro.
"""
import io
import json
import os
import sys
import uuid

import openpyxl

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')

import almacen_bd as bd  # noqa: E402
import encuestas_bd as ebd  # noqa: E402
import encuestas_hoja as hoja_mod  # noqa: E402
import formularios_libro as flibro  # noqa: E402
import nucleo_archivos as nucleo  # noqa: E402
from encuestas_modelo import formulario_vacio  # noqa: E402

USUARIO = 14
CARPETA = '/PRUEBAS FORMULARIOS'
LIBRO_NOMBRE = 'Libro con formularios.xlsx'
LIBRO = CARPETA + '/' + LIBRO_NOMBRE
FISICO = '/mnt/almacen/%d/archivos' % USUARIO


def crear_libro():
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Datos'
    hoja['A1'] = 'Presupuesto'
    hoja['A2'] = 1200
    hoja['A3'] = 800
    hoja['A4'] = '=A2+A3'
    memoria = io.BytesIO()
    libro.save(memoria)
    libro.close()
    memoria.seek(0)
    nucleo.subir(USUARIO, CARPETA, LIBRO_NOMBRE, memoria)
    bd.ejecutar('DELETE FROM vinculos_datos WHERE destino_ruta = %s', (LIBRO,))
    print('libro creado, con la fórmula A4 = A2+A3 (debe valer 2000)')


def crear_formulario(titulo):
    """Los mismos pasos del endpoint."""
    definicion = formulario_vacio(titulo)
    contenido = json.dumps(definicion, ensure_ascii=False, indent=2).encode('utf-8')
    nucleo.subir(USUARIO, CARPETA, titulo + '.forma', io.BytesIO(contenido))
    ruta_forma = CARPETA + '/' + titulo + '.forma'

    encuesta_id = definicion.get('id') or uuid.uuid4().hex
    definicion['id'] = encuesta_id
    ebd.registrar(encuesta_id, USUARIO, ruta_forma, titulo)
    fila = ebd.obtener(encuesta_id)

    respuestas = hoja_mod.construir(fila, definicion)
    if respuestas is None:
        raise RuntimeError('no se pudo construir la hoja de respuestas')
    nombre_respuestas = titulo[:80] + ' (respuestas).xlsx'
    nucleo.subir(USUARIO, CARPETA, nombre_respuestas, respuestas)
    hoja_mod.vincular(encuesta_id, CARPETA + '/' + nombre_respuestas)

    hoja = flibro.enlazar(USUARIO, encuesta_id, LIBRO)
    print('  «%s» → hoja «%s»' % (titulo, hoja))
    return hoja


def main():
    crear_libro()
    print('creando dos formularios desde el mismo libro:')
    crear_formulario('Formulario de Libro con formularios')
    crear_formulario('Formulario de Libro con formularios 2')

    libro = openpyxl.load_workbook(FISICO + LIBRO, data_only=True)
    print('\nhojas del libro:', libro.sheetnames)
    print('fórmula propia A4 (debe valer 2000):', libro['Datos']['A4'].value)
    libro.close()

    print('\nhojas que reciben respuestas:')
    for enlace in flibro.hojas_enlazadas(LIBRO):
        print('  %-14s ←  %s' % (enlace['destino_hoja'],
                                 enlace['origen_ruta'].split('/')[-1]))

    print('\narchivos creados en la carpeta:')
    for nombre in sorted(os.listdir(FISICO + CARPETA)):
        print('   ', nombre)


if __name__ == '__main__':
    main()
