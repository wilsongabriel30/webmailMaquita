#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Comprueba que al aligerar los archivos no se perdió ningún dato.

Pregunta de Wilson: si el archivo de Maquita pesa mucho menos que el de Google,
¿no faltará algo? ¿No estarán corruptos, o dejarán de abrirse o descargarse?

Se compara, hoja por hoja y celda por celda, el archivo TAL CUAL LO BAJA GOOGLE
contra el publicado en el Drive Maquita:

  · cuántas celdas CON DATO tiene cada uno (las vacías no cuentan);
  · si algún valor cambió;
  · si algún valor desapareció;
  · si el archivo publicado es un zip íntegro (que abra y se descargue bien).

Lo que se elimina al aligerar son celdas SIN contenido que solo llevaban un
estilo. Si la cuenta de celdas con dato coincide y ningún valor difiere, no se
perdió información.

Uso:  verificar_integridad.py <carpeta-crudo> <carpeta-publicada> [maximo]
"""
import os
import sys
import zipfile

import openpyxl


def valores(ruta):
    """{(hoja, fila, columna): valor} de todas las celdas CON dato."""
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    datos = {}
    try:
        for nombre in libro.sheetnames:
            pagina = libro[nombre]
            for fila in pagina.iter_rows():
                for celda in fila:
                    if celda.value is not None:
                        datos[(nombre, celda.row, celda.column)] = celda.value
    finally:
        libro.close()
    return datos


def main():
    crudo = sys.argv[1]
    publicado = sys.argv[2]
    maximo = int(sys.argv[3]) if len(sys.argv) > 3 else 40

    revisados = perdidos_tot = cambiados_tot = corruptos = 0
    problemas = []

    for carpeta, _, archivos in os.walk(crudo):
        for nombre in sorted(archivos):
            if not nombre.lower().endswith('.xlsx') or revisados >= maximo:
                continue
            rel = os.path.relpath(os.path.join(carpeta, nombre), crudo)
            destino = os.path.join(publicado, rel)
            if not os.path.exists(destino):
                continue

            # ¿El archivo publicado abre y se puede descargar entero?
            try:
                with zipfile.ZipFile(destino) as paquete:
                    if paquete.testzip() is not None:
                        raise zipfile.BadZipFile('contenido dañado')
            except Exception as excepcion:
                corruptos += 1
                problemas.append('CORRUPTO  %s (%s)' % (rel[:70], excepcion))
                continue

            try:
                antes = valores(os.path.join(crudo, rel))
                despues = valores(destino)
            except Exception as excepcion:
                problemas.append('ILEGIBLE  %s (%s)' % (rel[:70], excepcion))
                continue

            revisados += 1
            perdidos = [k for k in antes if k not in despues]
            cambiados = [k for k in antes
                         if k in despues and antes[k] != despues[k]]
            perdidos_tot += len(perdidos)
            cambiados_tot += len(cambiados)

            marca = 'OK ' if not (perdidos or cambiados) else '!! '
            print('%s %-62s datos: %7d -> %7d  perdidos:%5d  cambiados:%5d'
                  % (marca, rel[-62:], len(antes), len(despues),
                     len(perdidos), len(cambiados)))
            if perdidos[:2] or cambiados[:2]:
                for clave in (perdidos[:2] + cambiados[:2]):
                    print('       hoja %s fila %s col %s: Google=%r Maquita=%r'
                          % (clave[0], clave[1], clave[2],
                             antes.get(clave), despues.get(clave)))

    print()
    print('=' * 64)
    print('archivos comparados          : %d' % revisados)
    print('archivos corruptos           : %d' % corruptos)
    print('VALORES PERDIDOS en total    : %d' % perdidos_tot)
    print('VALORES CAMBIADOS en total   : %d' % cambiados_tot)
    for linea in problemas[:10]:
        print('  ' + linea)


if __name__ == '__main__':
    main()
