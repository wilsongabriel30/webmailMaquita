# -*- coding: utf-8 -*-
"""Apila los rangos de varias hojas de cálculo en un bloque de un destino.

Es el reemplazo de `QUERY({IMPORTRANGE(...);IMPORTRANGE(...)})` de Google:
`{a;b}` apilaba verticalmente y `QUERY` recortaba. Aquí se apila leyendo los
VALORES calculados de cada fuente y se escriben como valores en el destino,
igual que hace `api_vinculos.py` con un solo origen: sin fórmulas externas no
hay aviso de «actualizar vínculos» al abrir, y el dato siempre está fresco.
"""
import logging
import os

import openpyxl
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple

log = logging.getLogger(__name__)


ERRORES_EXCEL = ('#DIV/0!', '#N/A', '#VALUE!', '#REF!', '#NAME?', '#NUM!', '#NULL!',
                 '#¡DIV/0!', '#N/D', '#¡VALOR!', '#¡REF!', '#¿NOMBRE?', '#¡NUM!', '#¡NULO!')


def _es_error(valor):
    return isinstance(valor, str) and valor.strip().upper() in ERRORES_EXCEL


class LectorFuentes:
    """Lee rangos de varias matrices reutilizando cada libro abierto.

    Dos detalles que deciden si esto tarda segundos o media hora:
      - Se abre CADA archivo una sola vez, aunque de él salgan varias hojas
        (Esmeraldas aporta cuatro proyectos, por ejemplo).
      - Se lee con `iter_rows(values_only=True)`. En modo `read_only`, pedir
        celda por celda con `cell(row, column)` recorre la hoja en cada
        llamada: con los rangos anchos de meses (70 columnas × cientos de
        filas × 35 fuentes) eso no termina nunca.
    """

    def __init__(self):
        self._libros = {}

    def _libro(self, ruta):
        if ruta not in self._libros:
            self._libros[ruta] = openpyxl.load_workbook(
                ruta, data_only=True, read_only=True)
        return self._libros[ruta]

    def leer_rango(self, ruta_fisica_archivo, hoja, rango):
        """Valores calculados de un rango. Devuelve [] si la hoja no existe."""
        libro = self._libro(ruta_fisica_archivo)
        if hoja not in libro.sheetnames:
            log.warning('la hoja «%s» no existe en %s', hoja,
                        os.path.basename(ruta_fisica_archivo))
            return []
        col_min, fila_min, col_max, fila_max = range_boundaries(rango)
        pagina = libro[hoja]
        filas = pagina.iter_rows(min_row=fila_min, max_row=fila_max,
                                 min_col=col_min, max_col=col_max,
                                 values_only=True)
        # Los errores de las fórmulas de la matriz (`=Q5/P5` con lo planificado
        # vacío → #DIV/0!) llegaban al consolidado como texto «#DIV/0!» y lo
        # llenaban de errores que no son suyos. Van como celda vacía (11/09/2026).
        return [[None if _es_error(c) else c for c in fila] for fila in filas]

    def cerrar(self):
        for libro in self._libros.values():
            try:
                libro.close()
            except Exception:                      # cerrar nunca debe romper la corrida
                pass
        self._libros.clear()


def apilar(bloques):
    """Concatena verticalmente las matrices, rellenando al ancho mayor.

    Google apilaba rangos de anchos distintos sin protestar; openpyxl escribe
    fila a fila, así que se normaliza el ancho para que las columnas no se
    desalineen cuando una fuente trae menos columnas que otra.
    """
    matriz = [fila for bloque in bloques for fila in bloque]
    if not matriz:
        return []
    ancho = max(len(fila) for fila in matriz)
    return [fila + [None] * (ancho - len(fila)) for fila in matriz]


def abrir_destino(ruta_fisica_destino, hoja):
    """Abre el consolidado conservando sus fórmulas y formato."""
    libro = openpyxl.load_workbook(ruta_fisica_destino)
    pagina = libro[hoja] if hoja in libro.sheetnames else libro.create_sheet(hoja)
    return libro, pagina


def pegar(pagina, celda, matriz, filas_a_limpiar=0):
    """Escribe la matriz como valores a partir de `celda`.

    `filas_a_limpiar` sobrescribe con vacío ese número de filas, para que un
    consolidado que encoge no deje colgando los restos de la corrida anterior.
    """
    fila0, col0 = coordinate_to_tuple(celda)
    ancho = max((len(f) for f in matriz), default=0)
    # `cell(..., value=None)` NO borra la celda en openpyxl (deja lo que había):
    # las filas sobrantes de la corrida anterior y los errores heredados se
    # quedaban. Se asigna `.value` siempre, también cuando es None (11/09/2026).
    for i in range(max(len(matriz), filas_a_limpiar)):
        origen = matriz[i] if i < len(matriz) else []
        for j in range(ancho):
            pagina.cell(row=fila0 + i, column=col0 + j).value = (
                origen[j] if j < len(origen) else None)
