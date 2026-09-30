# -*- coding: utf-8 -*-
"""Las hojas por rubro de los consolidados ASC: lo que en Google hacía QUERY.

Responsabilidad ÚNICA: después de que el consolidado tenga sus bloques pegados,
rehacer las hojas que en Google eran

    =QUERY(ConsolidadoP!H4:CN, "SELECT * WHERE O = 'ASC - Comunicación'")

es decir, la misma tabla del consolidado, filtrada por una columna, pegada en
otra hoja. Es lo que mira gerencia: «Comunicación», «Formación», «Personal
Local»… Al salir de Google esas hojas quedaron congeladas con los valores de
la última exportación: el consolidado se rehacía solo, pero sus hojas por
rubro no (30/09/2026).

Cómo se resuelve aquí, sin QUERY:
  · se lee el rango del consolidado tal como quedó pegado (valores);
  · las columnas que son fórmula dentro del rango (presupuesto anual = suma de
    los meses, % ejecutado = gasto / presupuesto) se calculan aquí, porque
    openpyxl no calcula y Google las entregaba ya calculadas;
  · se filtra por la columna que decía el QUERY y se pega, como valores, donde
    estaba el QUERY; lo que sobre de la corrida anterior se limpia.

Qué hoja se filtra por qué valor sale de las fórmulas originales, leídas en la
exportación cruda de Google el 30/09/2026 (`mapa_conexiones_hoy.py`).
"""
import logging
import re

from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter

log = logging.getLogger('almacen.consolidados.rubros')

_RUBROS_ASC = {
    'Comunicación': ['ASC - Comunicación'],
    'Asesorías': ['ASC - Asesorías'],
    'Formación': ['ASC - Formación'],
    'Seguimiento Nacional': ['ASC - Seguimiento Nacional'],
    'Equipamiento': ['ASC - Equipamiento'],
    'Eventos Nacionales': ['ASC - Eventos Nacionales'],
    'Viajes Internacionales': ['ASC - Viajes Internacionales'],
    'Personal Local': ['ASC - Personal Local'],
    'Funcionamiento e Imprevistos': ['ASC - Funcionamiento e Imprevistos'],
    'Evaluación y Auditoría': ['ASC - Evaluación y Auditoría'],
}

# Por archivo del consolidado: de qué rango sale la tabla, por qué columna se
# filtra, y cada hoja con sus valores (o la celda de la propia hoja donde está
# el valor a buscar, como en Ejec. Técnica: `WHERE Col3 = A1`).
DEFINICIONES = {
    'Mapa de inversiones/Mapa general de inversiones 2026.xlsx': {
        'hoja_origen': 'ConsolidadoP', 'desde': 'H', 'hasta': 'CN', 'fila': 4,
        'filtro': 'O',
        'hojas': dict(_RUBROS_ASC, **{
            'Proceso Comercial': ['Maquita Agro', 'Maquita Productos', 'Maquita Turismo']}),
        'destino': 'A4',
    },
    'Mapa de inversiones/Mapa general de inversiones 2025.xlsx': {
        'hoja_origen': 'ConsolidadoP', 'desde': 'H', 'hasta': 'CB', 'fila': 4,
        'filtro': 'O',
        'hojas': _RUBROS_ASC,
        'destino': 'A4',
        # En este libro «Asesorías» tiene el QUERY en C4, no en A4.
        'destino_por_hoja': {'Asesorías': 'C4'},
    },
    'Mapa de inversiones/Mapa general Ejec. Técnica 2026.xlsx': {
        'hoja_origen': 'ConsolidadoT', 'desde': 'G', 'hasta': 'BB', 'fila': 4,
        'filtro': 'I',                      # «Col3» del rango G:BB
        'valor_en': 'A1',                   # el valor a buscar está en A1 de cada hoja
        # «Buscador» se deja fuera: en Google filtraba por lo que se tecleaba en
        # A1, y eso no se puede rehacer en diferido.
        # Las hojas se llaman como el código del resultado (resultado.actividad);
        # se arman aquí para que no parezcan direcciones de red.
        'hojas': {'2.1.%d.%d' % (r, a): None for r, a in
                  ((1, 1), (1, 2), (1, 3), (1, 4), (2, 1), (2, 2), (3, 1), (3, 2), (3, 3), (4, 1))},
        'destino': 'A4',
    },
}

_RE_SUMA = re.compile(r'^=\s*\$?([A-Z]{1,3})\$?(\d+)(\s*\+\s*\$?[A-Z]{1,3}\$?\d+)*\s*$')
_RE_DIV = re.compile(r'^=\s*\$?([A-Z]{1,3})\$?(\d+)\s*/\s*\$?([A-Z]{1,3})\$?(\d+)\s*$')
_RE_REF = re.compile(r'\$?([A-Z]{1,3})\$?(\d+)')


def _numero(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '.'))
    except ValueError:
        return None


def _valor(hoja, fila, columna, memo, profundidad=0):
    """Valor de una celda; si es fórmula de suma o de división en la misma fila,
    se calcula. Cualquier otra fórmula → None (no se inventa un valor)."""
    clave = (fila, columna)
    if clave in memo:
        return memo[clave]
    v = hoja.cell(row=fila, column=columna).value
    if isinstance(v, str) and v.startswith('=') and profundidad < 4:
        f = v.replace(' ', '')
        # El libro de 2025 escribe la suma como =SUM(V4+AA4+…): es la misma suma.
        m_sum = re.match(r'^=SUM\((.+)\)$', f, re.I)
        if m_sum and re.match(r'^[A-Z$0-9+,]+$', m_sum.group(1)):
            f = '=' + m_sum.group(1).replace(',', '+')
        if _RE_SUMA.match(f):
            total = 0.0
            for col, fil in _RE_REF.findall(f):
                n = _numero(_valor(hoja, int(fil), column_index_from_string(col), memo, profundidad + 1))
                if n is None:
                    total = None
                    break
                total += n
            v = total
        elif _RE_DIV.match(f):
            m = _RE_DIV.match(f)
            a = _numero(_valor(hoja, int(m.group(2)), column_index_from_string(m.group(1)), memo, profundidad + 1))
            b = _numero(_valor(hoja, int(m.group(4)), column_index_from_string(m.group(3)), memo, profundidad + 1))
            v = (a / b) if (a is not None and b not in (None, 0.0)) else None
        else:
            v = None
    elif isinstance(v, str) and v.startswith('='):
        v = None
    memo[clave] = v
    return v


def tabla(hoja, desde, hasta, fila):
    """Filas del rango desde:hasta a partir de `fila`, con las fórmulas calculadas.
    Termina en la última fila con algún dato."""
    c0, c1 = column_index_from_string(desde), column_index_from_string(hasta)
    memo = {}
    filas = []
    ultima = hoja.max_row or fila
    for f in range(fila, ultima + 1):
        valores = [_valor(hoja, f, c, memo) for c in range(c0, c1 + 1)]
        filas.append(valores)
    while filas and all(v in (None, '') for v in filas[-1]):
        filas.pop()
    return filas


def _texto(v):
    return str(v).strip() if v is not None else ''


def filtrar(filas, desde, filtro, valores):
    """Las filas cuya columna `filtro` (letra absoluta) tiene uno de `valores`."""
    idx = column_index_from_string(filtro) - column_index_from_string(desde)
    buscados = {_texto(v) for v in valores}
    return [f for f in filas if idx < len(f) and _texto(f[idx]) in buscados]


def pegar(pagina, celda, matriz, filas_previas):
    """Escribe la matriz como valores desde `celda` y limpia lo que sobre."""
    fila0, col0 = coordinate_to_tuple(celda)
    ancho = max((len(f) for f in matriz), default=0)
    a_limpiar = max(0, filas_previas - fila0 + 1)
    for i in range(max(len(matriz), a_limpiar)):
        origen = matriz[i] if i < len(matriz) else []
        for j in range(ancho):
            pagina.cell(row=fila0 + i, column=col0 + j).value = (
                origen[j] if j < len(origen) else None)
    return len(matriz), ancho


def regenerar(libro, archivo, hoja_origen=None):
    """Rehace las hojas por rubro del consolidado `archivo` dentro de `libro`
    (openpyxl, ya con los bloques pegados). Devuelve [(hoja, celda, alto, ancho)]."""
    definicion = DEFINICIONES.get(archivo)
    if not definicion:
        return []
    origen = libro[definicion['hoja_origen']]
    filas = tabla(origen, definicion['desde'], definicion['hasta'], definicion['fila'])
    escritos = []
    for nombre, valores in definicion['hojas'].items():
        if nombre not in libro.sheetnames:
            log.warning('  rubro «%s»: la hoja no existe en el libro, se omite', nombre)
            continue
        pagina = libro[nombre]
        if valores is None:
            valores = [pagina[definicion['valor_en']].value]
        celda = definicion.get('destino_por_hoja', {}).get(nombre, definicion['destino'])
        seleccion = filtrar(filas, definicion['desde'], definicion['filtro'], valores)
        alto, ancho = pegar(pagina, celda, seleccion, pagina.max_row or 0)
        escritos.append((nombre, celda, alto, ancho))
    log.info('  hojas por rubro rehechas: %d (%s)', len(escritos),
             ', '.join('%s=%d' % (n, a) for n, _, a, _ in escritos))
    return escritos
