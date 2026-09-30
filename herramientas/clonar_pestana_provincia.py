# -*- coding: utf-8 -*-
"""Clona una pestaña de provincia con TODO su formato y la adapta a otra provincia.

Rehacer una pestaña escribiendo celda a celda con la Api deja los datos bien pero
la hoja pelada: sin tabla, sin celdas combinadas, sin la imagen de la cabecera y
sin estilos (le pasó a BOLIVAR el 22/09/2026). El editor tampoco sabe duplicar
hojas (`hoja.Copy` no existe).

Aquí se copia la hoja DENTRO del paquete del `.xlsx`: su XML, su tabla, su dibujo
(reutilizando la misma imagen, sin duplicarla) y sus relaciones. Después se
adapta lo que distingue a una provincia de otra:

  · el rótulo «PROVINCIA : …» de la cabecera;
  · el encabezado «Territorio - …», en la celda y en la definición de la tabla;
  · la celda de servicio CB4, que es como escribe la provincia el formulario;
  · las fórmulas de la columna de territorio, que pasan a mirar la columna que
    le toca en «Prueba IFO».

Los valores en caché de las fórmulas se borran: los calcula el editor al abrir
(quien llame puede además pasar el libro por `encuestas_hoja_recalculo`).

Uso:
    python clonar_pestana_provincia.py <usuario> "<ruta>" <ORIGEN> <DESTINO> \
        --clave "Bolivar" --rotulo "BOLÍVAR" [--aplicar]

Autoría: Equipo de Tecnología Maquita — 2026-09-22
"""
import io
import posixpath
import re
import sys
import zipfile

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
sys.path.insert(0, '/home/sistemas/almacen-maquita/herramientas')

import openpyxl

import nucleo_archivos as nucleo
from encuestas_hoja_xml import RE_CELDA, RE_FILA, _celda_valor, _col_letras, _col_num
from formulas_provincias import (HOJA_BASE, FILA_CABECERAS_BASE, FILA_CABECERAS_PESTANA,
                                 plano)

FILA_ROTULO = 4
CELDA_CLAVE = 'CB4'


def _partes(z):
    """{nombre visible: (parte, rels de la parte, r:id)} de cada hoja."""
    libro = z.read('xl/workbook.xml').decode('utf-8')
    rels = {}
    for rel in re.findall(r'<Relationship\b[^>]*>', z.read('xl/_rels/workbook.xml.rels').decode('utf-8')):
        i = re.search(r'Id="([^"]+)"', rel)
        t = re.search(r'Target="([^"]+)"', rel)
        if i and t:
            rels[i.group(1)] = 'xl/' + t.group(1).lstrip('/')
    salida = {}
    for etiqueta in re.findall(r'<sheet\b[^>]*/?>', libro):
        nombre = re.search(r'name="([^"]*)"', etiqueta)
        rid = re.search(r'r:id="([^"]*)"', etiqueta)
        if nombre and rid:
            parte = rels[rid.group(1)]
            salida[nombre.group(1)] = (parte, etiqueta, rid.group(1))
    return salida


def _libre(z, patron, plantilla):
    usados = {int(m.group(1)) for p in z.namelist()
              for m in [re.fullmatch(patron, p)] if m}
    n = 1
    while n in usados:
        n += 1
    return plantilla % n


def _columna_de(hoja, titulo_normalizado):
    for c in range(1, 120):
        valor = hoja.cell(row=FILA_CABECERAS_PESTANA, column=c).value
        if valor and plano(valor).startswith(titulo_normalizado):
            return c
    return None


def clonar(contenido, origen, destino, clave, rotulo):
    z = zipfile.ZipFile(io.BytesIO(contenido))
    hojas = _partes(z)
    if origen not in hojas:
        raise SystemExit('no existe la pestaña «%s»' % origen)
    parte_origen, etiqueta_origen, rid_origen = hojas[origen]
    parte_vieja = hojas.get(destino, (None, None, None))[0]

    libro_py = openpyxl.load_workbook(io.BytesIO(contenido))
    modelo = libro_py[origen]
    base = libro_py[HOJA_BASE]
    col_territorio_pestana = _columna_de(modelo, 'TERRITORIO -')
    titulo_territorio = 'Territorio - %s' % destino
    col_base_destino = col_base_origen = None
    for c in range(1, 120):
        valor = base.cell(row=FILA_CABECERAS_BASE, column=c).value
        if not valor:
            continue
        if plano(valor) == plano(titulo_territorio):
            col_base_destino = c
        if plano(valor) == plano('Territorio - %s' % origen):
            col_base_origen = c
    if not (col_territorio_pestana and col_base_destino):
        raise SystemExit('no encuentro la columna de territorio (pestaña %s, base %s)'
                         % (col_territorio_pestana, col_base_destino))

    # Piezas nuevas del paquete.
    nueva_hoja = _libre(z, r'xl/worksheets/sheet(\d+)\.xml', 'xl/worksheets/sheet%d.xml')
    nuevo_dibujo = _libre(z, r'xl/drawings/drawing(\d+)\.xml', 'xl/drawings/drawing%d.xml')
    nueva_tabla = _libre(z, r'xl/tables/table(\d+)\.xml', 'xl/tables/table%d.xml')

    rels_origen_ruta = posixpath.join(posixpath.dirname(parte_origen), '_rels',
                                      posixpath.basename(parte_origen) + '.rels')
    rels_origen = z.read(rels_origen_ruta).decode('utf-8') if rels_origen_ruta in z.namelist() else ''
    tabla_origen = dibujo_origen = None
    for objetivo in re.findall(r'Target="([^"]+)"', rels_origen):
        completo = posixpath.normpath(posixpath.join(posixpath.dirname(parte_origen), objetivo))
        if '/tables/' in completo:
            tabla_origen = completo
        elif '/drawings/' in completo and completo.endswith('.xml'):
            dibujo_origen = completo

    piezas = {}
    nombre_tabla = 'Resp_%s' % re.sub(r'[^A-Za-z0-9_]', '_', plano(destino))
    tabla_vieja = None
    if tabla_origen:
        tabla_vieja = re.search(r'displayName="([^"]*)"',
                                z.read(tabla_origen).decode('utf-8')).group(1)
    # ── la hoja ──────────────────────────────────────────────────────────
    xml = z.read(parte_origen).decode('utf-8')
    letra_destino = _col_letras(col_base_destino)
    letra_origen = _col_letras(col_base_origen) if col_base_origen else None

    def arregla_fila(m):
        fila = int(m.group(1))
        interior = m.group(3) or ''

        def arregla_celda(c):
            columna, coord = _col_num(c.group(1)), c.group(1) + c.group(2)
            atributos, dentro = c.group(3) or '', c.group(4) or ''
            estilo = re.search(r's="(\d+)"', atributos)
            estilo = estilo.group(1) if estilo else None
            if coord == CELDA_CLAVE:
                return _celda_valor(coord, estilo, clave)
            if fila == FILA_ROTULO and plano(origen) in plano(_texto(dentro, atributos, z)):
                return _celda_valor(coord, estilo, rotulo)
            if fila == FILA_CABECERAS_PESTANA and columna == col_territorio_pestana:
                return _celda_valor(coord, estilo, titulo_territorio)
            if '<f' in dentro:
                formula = re.search(r'<f[^>]*>(.*?)</f>', dentro, re.S)
                texto = formula.group(1) if formula else ''
                if columna == col_territorio_pestana and letra_origen:
                    # En el XML la fórmula va escapada (&apos;), así que se
                    # cambia la LETRA de columna allí donde una referencia la
                    # nombra: tras «!» (inicio del rango) o tras «:» (final).
                    texto = re.sub(r'(?<=[!:])\$?%s(\$?\d)' % letra_origen,
                                   lambda x: '%s%s' % (letra_destino, x.group(1)), texto)
                # Sin el valor en caché: lo calcula el editor al abrir.
                return '<c r="%s"%s><f>%s</f></c>' % (coord, atributos.replace(' t="str"', ''), texto)
            return c.group(0)
        return '<row r="%s"%s>%s</row>' % (m.group(1), m.group(2),
                                           RE_CELDA.sub(arregla_celda, interior))

    xml = RE_FILA.sub(arregla_fila, xml)
    if tabla_vieja:
        # Las fórmulas propias (TOTAL, DURACIÓN) nombran la tabla de la hoja:
        # sin esto, la copia seguiría sumando las filas de la ORIGINAL.
        xml = xml.replace('%s[' % tabla_vieja, '%s[' % nombre_tabla)
    piezas[nueva_hoja] = xml
    if rels_origen:
        nuevo_rels = rels_origen
        if tabla_origen:
            nuevo_rels = nuevo_rels.replace(posixpath.basename(tabla_origen),
                                            posixpath.basename(nueva_tabla))
        if dibujo_origen:
            nuevo_rels = nuevo_rels.replace(posixpath.basename(dibujo_origen),
                                            posixpath.basename(nuevo_dibujo))
        piezas[posixpath.join(posixpath.dirname(nueva_hoja), '_rels',
                              posixpath.basename(nueva_hoja) + '.rels')] = nuevo_rels
    # ── la tabla ─────────────────────────────────────────────────────────
    if tabla_origen:
        xml_tabla = z.read(tabla_origen).decode('utf-8')
        ids = [int(x) for x in re.findall(r'<table\b[^>]*\bid="(\d+)"', xml_tabla)] or [1]
        xml_tabla = re.sub(r'(<table\b[^>]*\bid=")\d+(")',
                           lambda m: m.group(1) + str(max(ids) + 50) + m.group(2), xml_tabla)
        if tabla_vieja:
            xml_tabla = xml_tabla.replace('%s[' % tabla_vieja, '%s[' % nombre_tabla)
        xml_tabla = re.sub(r'\bname="[^"]*"', 'name="%s"' % nombre_tabla, xml_tabla, count=1)
        xml_tabla = re.sub(r'\bdisplayName="[^"]*"', 'displayName="%s"' % nombre_tabla,
                           xml_tabla, count=1)
        # El encabezado de la columna de territorio, también en la tabla.
        inicio = _col_num(re.search(r'\bref="([A-Z]+)\d+', xml_tabla).group(1))
        columnas = re.findall(r'<tableColumn\b[^>]*?\bname="([^"]*)"', xml_tabla)
        posicion = col_territorio_pestana - inicio
        if 0 <= posicion < len(columnas):
            viejo = columnas[posicion]
            xml_tabla = xml_tabla.replace('name="%s"' % viejo, 'name="%s"' % titulo_territorio, 1)
        piezas[nueva_tabla] = xml_tabla
    # ── el dibujo (misma imagen) ─────────────────────────────────────────
    if dibujo_origen:
        piezas[nuevo_dibujo] = z.read(dibujo_origen).decode('utf-8')
        rels_dibujo = posixpath.join(posixpath.dirname(dibujo_origen), '_rels',
                                     posixpath.basename(dibujo_origen) + '.rels')
        if rels_dibujo in z.namelist():
            piezas[posixpath.join(posixpath.dirname(nuevo_dibujo), '_rels',
                                  posixpath.basename(nuevo_dibujo) + '.rels')] = \
                z.read(rels_dibujo).decode('utf-8')
    # ── workbook, relaciones y tipos ─────────────────────────────────────
    rels_libro = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    usados = {int(m) for m in re.findall(r'Id="rId(\d+)"', rels_libro)}
    nuevo_id = 'rId%d' % (max(usados) + 1)
    rels_libro = rels_libro.replace('</Relationships>',
        '<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/'
        '2006/relationships/worksheet" Target="%s"/></Relationships>'
        % (nuevo_id, nueva_hoja[len('xl/'):]))
    libro_xml = z.read('xl/workbook.xml').decode('utf-8')
    if parte_vieja:                       # la hoja pelada se sustituye en su sitio
        etiqueta_vieja = _partes(z)[destino][1]
        libro_xml = libro_xml.replace(
            etiqueta_vieja, re.sub(r'r:id="[^"]*"', 'r:id="%s"' % nuevo_id,
                                   re.sub(r'\bsheetId="\d+"', 'sheetId="%d"' % (max(
                                       int(x) for x in re.findall(r'sheetId="(\d+)"', libro_xml)) + 1),
                                       etiqueta_vieja)))
    tipos = z.read('[Content_Types].xml').decode('utf-8')
    anadir = ['<Override PartName="/%s" ContentType="application/vnd.openxmlformats-'
              'officedocument.spreadsheetml.worksheet+xml"/>' % nueva_hoja]
    if tabla_origen:
        anadir.append('<Override PartName="/%s" ContentType="application/vnd.openxmlformats-'
                      'officedocument.spreadsheetml.table+xml"/>' % nueva_tabla)
    if dibujo_origen:
        anadir.append('<Override PartName="/%s" ContentType="application/vnd.openxmlformats-'
                      'officedocument.drawing+xml"/>' % nuevo_dibujo)
    tipos = tipos.replace('</Types>', ''.join(anadir) + '</Types>')

    piezas['xl/_rels/workbook.xml.rels'] = rels_libro
    piezas['xl/workbook.xml'] = libro_xml
    piezas['[Content_Types].xml'] = tipos

    fuera = {parte_vieja} if parte_vieja else set()
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, 'w', zipfile.ZIP_DEFLATED) as nuevo:
        for elemento in z.infolist():
            if elemento.filename in fuera:
                continue
            datos = piezas.pop(elemento.filename).encode('utf-8') \
                if elemento.filename in piezas else z.read(elemento.filename)
            nuevo.writestr(elemento, datos)
        for nombre, datos in piezas.items():
            nuevo.writestr(nombre, datos.encode('utf-8'))
    salida.seek(0)
    return salida


def _texto(dentro, atributos, z):
    if 't="inlineStr"' in atributos:
        return ''.join(re.findall(r'<t[^>]*>(.*?)</t>', dentro, re.S))
    if 't="s"' in atributos:
        v = re.search(r'<v>(\d+)</v>', dentro)
        if v:
            cadenas = _texto.cache = getattr(_texto, 'cache', None) or [
                ''.join(re.findall(r'<t[^>]*>(.*?)</t>', si, re.S))
                for si in re.findall(r'<si>(.*?)</si>',
                                     z.read('xl/sharedStrings.xml').decode('utf-8'), re.S)]
            i = int(v.group(1))
            return cadenas[i] if i < len(cadenas) else ''
    v = re.search(r'<v>(.*?)</v>', dentro, re.S)
    return v.group(1) if v else ''


def principal(argumentos):
    usuario, ruta, origen, destino = int(argumentos[0]), argumentos[1], argumentos[2], argumentos[3]
    clave = argumentos[argumentos.index('--clave') + 1] if '--clave' in argumentos else destino
    rotulo = argumentos[argumentos.index('--rotulo') + 1] if '--rotulo' in argumentos else destino
    fisica = nucleo.ruta_fisica(usuario, ruta)
    salida = clonar(open(fisica, 'rb').read(), origen, destino, clave, rotulo)
    datos = salida.getvalue()
    open('/tmp/pestana-clonada.xlsx', 'wb').write(datos)
    print('  resultado en /tmp/pestana-clonada.xlsx (%.0f KB)' % (len(datos) / 1024))
    if '--aplicar' in argumentos:
        import sala_editor
        from api_onlyoffice import _base_documento, invalidar_cache
        dentro = sala_editor.usuarios_conectados(_base_documento(usuario, ruta))
        if dentro is None or dentro:
            print('  NO se aplica: el Excel está abierto en el editor (%s)' % dentro)
            return 2
        import encuestas_hoja_recalculo as recalculo
        carpeta, _, nombre = ruta.rpartition('/')
        salida.seek(0)
        nucleo.subir(usuario, carpeta or '/', nombre, salida)
        calculado = recalculo.recalcular(usuario, ruta)
        if calculado:
            nucleo.subir(usuario, carpeta or '/', nombre, io.BytesIO(calculado))
            print('  recalculado por el editor')
        invalidar_cache(usuario, ruta)
        print('  aplicado al Drive')
    return 0


if __name__ == '__main__':
    sys.exit(principal(sys.argv[1:]))
