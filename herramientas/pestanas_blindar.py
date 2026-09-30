# -*- coding: utf-8 -*-
"""Envuelve en SI.ERROR las fórmulas propias que fallan en las filas sin respuesta.

DURACIÓN resta dos fechas; en una fila preparada pero todavía vacía la resta da
`#¡VALOR!`. El resto de columnas propias ya iban blindadas desde el 21/09/2026;
estas se quedaron fuera al rehacer las pestañas (22/09/2026).

Trabaja sobre el paquete del `.xlsx`: cambia la fórmula y borra el valor en
caché, para que el editor lo recalcule al abrir.

Uso:
    python pestanas_blindar.py <usuario> "<ruta>" [--aplicar]

Autoría: Equipo de Tecnología Maquita — 2026-09-22
"""
import io
import re
import sys
import zipfile

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
sys.path.insert(0, '/home/sistemas/almacen-maquita/herramientas')

import openpyxl

import nucleo_archivos as nucleo
from encuestas_hoja_xml import RE_CELDA, RE_FILA, unescape
from formulas_provincias import HOJA_BASE, FILA_CABECERAS_PESTANA
from pestanas_encabezados import _hojas_del_paquete

PRIMERA_FILA = FILA_CABECERAS_PESTANA + 1
ULTIMA_FILA = FILA_CABECERAS_PESTANA + 100


def a_blindar(fisica):
    """{nombre de hoja: {coordenada: fórmula blindada}}."""
    libro = openpyxl.load_workbook(fisica)
    salida = {}
    for nombre in libro.sheetnames:
        if nombre == HOJA_BASE:
            continue
        hoja = libro[nombre]
        celdas = {}
        for fila in range(PRIMERA_FILA, ULTIMA_FILA + 1):
            for columna in range(1, 80):
                celda = hoja.cell(row=fila, column=columna)
                valor = celda.value
                if not (isinstance(valor, str) and valor.startswith('=')):
                    continue
                cuerpo = valor[1:]
                if cuerpo.upper().startswith(('IFERROR', 'SI.ERROR')):
                    continue
                if 'INDEX(' in cuerpo.upper() or 'AGGREGATE(' in cuerpo.upper():
                    continue                      # las del formulario ya lo llevan
                celdas[celda.coordinate] = '=IFERROR(%s,"")' % cuerpo
        if celdas:
            salida[nombre] = celdas
    return salida


def _aplicar_en_hoja(xml, celdas):
    def en_fila(m):
        if not (PRIMERA_FILA <= int(m.group(1)) <= ULTIMA_FILA):
            return m.group(0)

        def en_celda(c):
            coordenada = c.group(1) + c.group(2)
            if coordenada not in celdas:
                return c.group(0)
            interior = c.group(4) or ''
            if '<f' not in interior:
                return c.group(0)
            # Fórmula nueva y sin valor en caché: lo recalcula el editor.
            return '<c r="%s"%s><f>%s</f></c>' % (
                coordenada, c.group(3) or '',
                celdas[coordenada][1:].replace('&', '&amp;').replace('<', '&lt;'))
        return '<row r="%s"%s>%s</row>' % (m.group(1), m.group(2),
                                           RE_CELDA.sub(en_celda, m.group(3) or ''))
    return RE_FILA.sub(en_fila, xml)


def principal(usuario, ruta, aplicar):
    fisica = nucleo.ruta_fisica(usuario, ruta)
    pendientes = a_blindar(fisica)
    for nombre, celdas in pendientes.items():
        print('  %-12s %d fórmulas' % (nombre, len(celdas)))
    if not pendientes:
        print('  nada que blindar')
        return 0
    contenido = open(fisica, 'rb').read()
    entrada = zipfile.ZipFile(io.BytesIO(contenido))
    hojas = _hojas_del_paquete(entrada)
    partes = {}
    for nombre, celdas in pendientes.items():
        parte = hojas.get(nombre, (None, None))[0]
        if not parte:
            continue
        partes[parte] = _aplicar_en_hoja(entrada.read(parte).decode('utf-8'), celdas)
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, 'w', zipfile.ZIP_DEFLATED) as z:
        for elemento in entrada.infolist():
            datos = partes[elemento.filename].encode('utf-8') \
                if elemento.filename in partes else entrada.read(elemento.filename)
            z.writestr(elemento, datos)
    salida.seek(0)
    open('/tmp/pestanas-blindadas.xlsx', 'wb').write(salida.getvalue())
    print('  resultado en /tmp/pestanas-blindadas.xlsx')
    if aplicar:
        import sala_editor
        from api_onlyoffice import _base_documento, invalidar_cache
        dentro = sala_editor.usuarios_conectados(_base_documento(usuario, ruta))
        if dentro is None or dentro:
            print('  NO se aplica: el Excel está abierto en el editor (%s)' % dentro)
            return 2
        carpeta, _, nombre_archivo = ruta.rpartition('/')
        salida.seek(0)
        nucleo.subir(usuario, carpeta or '/', nombre_archivo, salida)
        invalidar_cache(usuario, ruta)
        print('  aplicado al Drive')
    return 0


if __name__ == '__main__':
    sys.exit(principal(int(sys.argv[1]), sys.argv[2], '--aplicar' in sys.argv))
