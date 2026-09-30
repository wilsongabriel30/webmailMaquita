# -*- coding: utf-8 -*-
"""Pone a las pestañas de provincia el encabezado que le toca a cada columna.

Por qué hace falta: al insertar columnas dentro de la tabla de una pestaña, el
nombre de la columna vive en DOS sitios — la celda de la fila 5 y la definición
de la tabla (`xl/tables/*.xml`)—. El editor solo hacía caso al segundo: escribir
la celda con la Api no bastaba, y al reabrir el libro restauraba el nombre
antiguo («¿Qué tipo de mecanismos…») o el genérico («Column1») que él mismo
había puesto (22/09/2026).

Esto trabaja sobre el paquete del `.xlsx` directamente, sin editor:

  · el nombre correcto de cada columna se deduce de SU PROPIA FÓRMULA: una
    columna que trae `INDEX('Prueba IFO'!AB$5:…)` se llama como la columna AB de
    la hoja base. Es decir, el título también sale de «Prueba IFO»;
  · las columnas propias de la persona (TOTAL, DURACIÓN) conservan su nombre;
  · se corrigen la celda de la fila 5 Y el `name=` de la tabla, y las fórmulas
    estructuradas que citaran el nombre viejo.

Uso:
    python pestanas_encabezados.py <usuario> "<ruta>" [--aplicar]

Autoría: Equipo de Tecnología Maquita — 2026-09-22
"""
import io
import posixpath
import re
import shutil
import sys
import zipfile

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
sys.path.insert(0, '/home/sistemas/almacen-maquita/herramientas')

import openpyxl

import nucleo_archivos as nucleo
from encuestas_hoja_xml import (_cadenas, _celda_valor, _col_letras, _col_num, _rels,
                                _renombrar_en_formulas, _renombrar_columna_tabla,
                                _texto_celda, RE_CELDA, RE_FILA)
from formulas_provincias import HOJA_BASE, FILA_CABECERAS_BASE, FILA_CABECERAS_PESTANA

RE_INDICE = re.compile(r"INDEX\('?" + re.escape(HOJA_BASE) + r"'?!\$?([A-Z]+)\$?\d+")


def nombres_correctos(libro, nombre_hoja):
    """{columna: nombre que debe tener} para una pestaña, según sus fórmulas."""
    base, hoja = libro[HOJA_BASE], libro[nombre_hoja]
    titulo_base = {c: str(base.cell(row=FILA_CABECERAS_BASE, column=c).value)
                   for c in range(1, 120)
                   if base.cell(row=FILA_CABECERAS_BASE, column=c).value}
    salida = {}
    for columna in range(1, 120):
        if not hoja.cell(row=FILA_CABECERAS_PESTANA, column=columna).value:
            continue
        for fila in range(FILA_CABECERAS_PESTANA + 1, FILA_CABECERAS_PESTANA + 20):
            valor = hoja.cell(row=fila, column=columna).value
            if isinstance(valor, str) and valor.startswith('='):
                encontrado = RE_INDICE.search(valor)
                if encontrado:
                    columna_base = _col_num(encontrado.group(1))
                    if columna_base in titulo_base:
                        salida[columna] = titulo_base[columna_base]
                break
    return salida


def _hojas_del_paquete(z):
    """{nombre visible: (parte de la hoja, parte de su tabla o None)}."""
    libro = z.read('xl/workbook.xml').decode('utf-8')
    rels = _rels(z, 'xl/workbook.xml')
    salida = {}
    for etiqueta in re.findall(r'<sheet\b[^>]*/?>', libro):
        nombre = re.search(r'name="([^"]*)"', etiqueta)
        rid = re.search(r'r:id="([^"]*)"', etiqueta)
        if not (nombre and rid and rid.group(1) in rels):
            continue
        parte = rels[rid.group(1)]
        tabla = None
        for objetivo in _rels(z, parte).values():
            if '/tables/' in objetivo and objetivo in z.namelist():
                tabla = objetivo
        salida[re.sub(r'&amp;', '&', nombre.group(1))] = (parte, tabla)
    return salida


def _reescribir_encabezado(xml_hoja, fila_cabecera, columna, cadenas, texto):
    """Deja la celda de encabezado con el texto dado, en línea."""
    def en_fila(m):
        if int(m.group(1)) != fila_cabecera:
            return m.group(0)
        interior = m.group(3) or ''

        def en_celda(c):
            if _col_num(c.group(1)) != columna or int(c.group(2)) != fila_cabecera:
                return c.group(0)
            estilo = re.search(r's="(\d+)"', c.group(3) or '')
            return _celda_valor('%s%d' % (_col_letras(columna), fila_cabecera),
                                estilo.group(1) if estilo else None, texto)
        return '<row r="%s"%s>%s</row>' % (m.group(1), m.group(2),
                                           RE_CELDA.sub(en_celda, interior))
    return RE_FILA.sub(en_fila, xml_hoja)


def principal(usuario, ruta, aplicar):
    fisica = nucleo.ruta_fisica(usuario, ruta)
    libro = openpyxl.load_workbook(fisica)
    correctos = {n: nombres_correctos(libro, n) for n in libro.sheetnames
                 if n != HOJA_BASE and libro[n].max_row > FILA_CABECERAS_PESTANA}

    contenido = open(fisica, 'rb').read()
    entrada = zipfile.ZipFile(io.BytesIO(contenido))
    hojas = _hojas_del_paquete(entrada)
    cadenas = _cadenas(entrada)
    cambios = {}
    partes = {}
    for nombre, esperados in correctos.items():
        parte, tabla = hojas.get(nombre, (None, None))
        if not parte or not tabla:
            print('  %-12s sin tabla: se deja' % nombre)
            continue
        xml_hoja = partes.get(parte, entrada.read(parte).decode('utf-8'))
        xml_tabla = partes.get(tabla, entrada.read(tabla).decode('utf-8'))
        actuales = re.findall(r'<tableColumn\b[^>]*?\bname="([^"]*)"', xml_tabla)
        inicio = None
        for m in re.finditer(r'\bref="([A-Z]+)\d+:[A-Z]+\d+"', xml_tabla):
            inicio = _col_num(m.group(1))
            break
        renombres = []
        for columna, texto in sorted(esperados.items()):
            posicion = columna - (inicio or 1)
            if not (0 <= posicion < len(actuales)):
                continue
            viejo = actuales[posicion]
            if viejo == texto:
                continue
            renombres.append((viejo, texto))
            xml_tabla = _renombrar_columna_tabla(xml_tabla, viejo, texto)
            xml_tabla = _renombrar_en_formulas(xml_tabla, viejo, texto)
            xml_hoja = _renombrar_en_formulas(xml_hoja, viejo, texto)
            xml_hoja = _reescribir_encabezado(xml_hoja, FILA_CABECERAS_PESTANA,
                                              columna, cadenas, texto)
        if renombres:
            partes[parte], partes[tabla] = xml_hoja, xml_tabla
            cambios[nombre] = renombres
            print('  %-12s %d encabezados corregidos' % (nombre, len(renombres)))
        else:
            print('  %-12s ya estaba bien' % nombre)

    if not partes:
        print('  nada que cambiar')
        return 0
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, 'w', zipfile.ZIP_DEFLATED) as z:
        for elemento in entrada.infolist():
            datos = partes[elemento.filename].encode('utf-8') \
                if elemento.filename in partes else entrada.read(elemento.filename)
            z.writestr(elemento, datos)
    salida.seek(0)
    open('/tmp/pestanas-encabezados.xlsx', 'wb').write(salida.getvalue())
    print('  resultado en /tmp/pestanas-encabezados.xlsx')
    if aplicar:
        import sala_editor
        from api_onlyoffice import _base_documento, invalidar_cache
        dentro = sala_editor.usuarios_conectados(_base_documento(usuario, ruta))
        if dentro is None or dentro:
            print('  NO se aplica: el Excel está abierto en el editor (%s)' % dentro)
            return 2
        carpeta, _, nombre_archivo = ruta.rpartition('/')
        salida.seek(0)
        nucleo.subir(usuario, carpeta or '/', nombre_archivo, salida)
        invalidar_cache(usuario, ruta)
        print('  aplicado al Drive')
    return 0


if __name__ == '__main__':
    sys.exit(principal(int(sys.argv[1]), sys.argv[2], '--aplicar' in sys.argv))
