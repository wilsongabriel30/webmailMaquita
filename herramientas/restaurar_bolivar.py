# -*- coding: utf-8 -*-
"""Reconstruye la pestaña BOLIVAR de «prueba IFO (respuestas).xlsx».

La pestaña se perdió el 21/09/2026 (entre las 16:06 y las 16:08 desapareció del
libro y volvió vacía). No se recuperan sus fórmulas antiguas porque apuntaban a
columnas de la hoja base que ya no están en el mismo sitio: se rehace a partir
de COTOPAXI, que es la pestaña que más se le parecía (5 diferencias de rótulo en
la versión del 21/09) y que ya está al día, y se le cambia lo suyo:

  · la cabecera «PROVINCIA : BOLÍVAR»;
  · el encabezado «Territorio - BOLIVAR», apuntando a su columna de la base;
  · la celda de servicio CB4 = «Bolivar», que es como el formulario escribe la
    provincia.

El editor no sabe duplicar hojas (`hoja.Copy` no existe), así que se escribe
celda a celda: cabecera, encabezados, celda de servicio y fórmulas. Los TOTAL y
la DURACIÓN se copian de COTOPAXI, traduciendo sus referencias de tabla
(`Resp_COTOPAXI_2[[#This Row],[…]]`) a referencias normales: la pestaña que se
rehace no tiene tabla y esos nombres darían `#¿NOMBRE?`.

Uso:
    python restaurar_bolivar.py <usuario> "<ruta>" [--aplicar]

Autoría: Equipo de Tecnología Maquita — 2026-09-22
"""
import io
import os
import re
import secrets
import sys
import time

sys.path.insert(0, '/home/sistemas/almacen-maquita/servicio')
sys.path.insert(0, '/home/sistemas/almacen-maquita/herramientas')

import openpyxl
import requests

import nucleo_archivos as nucleo
from api_onlyoffice import firmar_jwt, url_interna_ds, _reescribir_url_interna
from config_almacen import URL_PUBLICA
from encuestas_hoja_recalculo import CARPETA_SCRIPTS, URL_SCRIPTS
from formulas_provincias import (HOJA_BASE, FILA_CABECERAS_BASE, FILA_CABECERAS_PESTANA,
                                 letra, plano, _js)


def _js_texto(valor):
    """Como `_js`, pero además con los saltos de línea escapados.

    Las celdas de la cabecera institucional los traen, y un salto literal dentro
    de la cadena rompe el guion: el editor devolvía `{"error":-3}` sin más
    explicación (22/09/2026)."""
    return _js(str(valor).replace('\r\n', '\\n').replace('\n', '\\n')
               .replace('\r', '\\n').replace('\t', ' '))


RE_TABLA_RANGO = re.compile(r"Resp_[A-Za-z0-9_]+\[\[#This Row\],\[([^\]]+)\]:\[([^\]]+)\]\]")
RE_TABLA_UNA = re.compile(r"Resp_[A-Za-z0-9_]+\[\[#This Row\],\[([^\]]+)\]\]")


def sin_referencias_de_tabla(formula, columnas, fila):
    """La misma fórmula con referencias normales en vez de estructuradas.

    Las de la plantilla nombran su tabla (`Resp_COTOPAXI_2[[#This Row],[TOTAL
    SOCIOS]]`); la pestaña que se rehace no tiene tabla, así que esos nombres
    darían `#¿NOMBRE?`. `columnas` es {nombre de encabezado: número de columna}.
    """
    def una(m):
        columna = columnas.get(plano(m.group(1)))
        return '%s%d' % (letra(columna), fila) if columna else m.group(0)

    def rango(m):
        desde = columnas.get(plano(m.group(1)))
        hasta = columnas.get(plano(m.group(2)))
        if not (desde and hasta):
            return m.group(0)
        return '%s%d:%s%d' % (letra(desde), fila, letra(hasta), fila)
    return RE_TABLA_UNA.sub(una, RE_TABLA_RANGO.sub(rango, formula))


PLANTILLA = 'COTOPAXI'
DESTINO = 'BOLIVAR'
CLAVE = 'Bolivar'                 # como lo escribe el formulario
ROTULO = 'BOLÍVAR'                # como se lee en la cabecera
TERRITORIO = 'Territorio - BOLIVAR'
FILAS = 100


def guion(fisica):
    libro = openpyxl.load_workbook(fisica)
    base, modelo = libro[HOJA_BASE], libro[PLANTILLA]
    columnas_base = {}
    for c in range(1, 120):
        valor = base.cell(row=FILA_CABECERAS_BASE, column=c).value
        if valor:
            columnas_base.setdefault(plano(valor), c)
    col_territorio = columnas_base.get(plano(TERRITORIO))
    if not col_territorio:
        raise SystemExit('la hoja base no tiene la columna «%s»' % TERRITORIO)

    ancho = max(c for c in range(1, 120)
                if modelo.cell(row=FILA_CABECERAS_PESTANA, column=c).value)
    # Se escribe SOBRE la hoja vacía que quedó; borrarla y recrearla hacía
    # fallar al editor (error -3, 22/09/2026).
    lineas = ['var hoja = Api.GetSheet(%s);' % _js(DESTINO),
              'if (!hoja) { Api.AddSheet(%s); hoja = Api.GetSheet(%s); }'
              % (_js(DESTINO), _js(DESTINO)),
              'if (hoja) {']
    # Cabecera institucional (filas 1-4) y encabezados (fila 5), copiados de la
    # plantilla; el rótulo de la provincia es el de BOLÍVAR.
    for fila in range(1, FILA_CABECERAS_PESTANA + 1):
        for columna in range(1, ancho + 1):
            valor = modelo.cell(row=fila, column=columna).value
            if valor in (None, '') or (isinstance(valor, str) and valor.startswith('=')):
                continue
            texto = str(valor)
            if fila == 4 and plano(texto) in ('COTOPAXI',):
                texto = ROTULO
            if fila == FILA_CABECERAS_PESTANA and plano(texto).startswith('TERRITORIO -'):
                texto = TERRITORIO
            lineas.append('  hoja.GetRange("%s%d").SetValue(%s);'
                          % (letra(columna), fila, _js_texto(texto)))
    # Celda de servicio y columna de servicio, como en las demás pestañas.
    lineas.append('  hoja.GetRange("CB%d").SetValue(%s);'
                  % (FILA_CABECERAS_PESTANA - 1, _js(CLAVE)))
    lineas.append('  hoja.GetRange("CB%d").SetValue(%s);'
                  % (FILA_CABECERAS_PESTANA - 2,
                     _js('Provincia tal como la guarda el formulario. No borrar.')))

    # Cada columna: la del formulario trae su dato de la base; las propias
    # copian la fórmula que tiene COTOPAXI en su primera fila de datos.
    primera = FILA_CABECERAS_PESTANA + 1
    rango_filas = "'%s'!$A$%d:$A$%d" % (HOJA_BASE, FILA_CABECERAS_BASE + 1, 1005)
    col_prov = letra(columnas_base[plano('Provincia')])
    destino_de = {}
    for columna in range(1, ancho + 1):
        titulo = modelo.cell(row=FILA_CABECERAS_PESTANA, column=columna).value
        if not titulo:
            continue
        clave = plano(titulo)
        if clave.startswith('TERRITORIO -'):
            destino_de[columna] = col_territorio
        elif clave in columnas_base:
            destino_de[columna] = columnas_base[clave]
        elif plano('Hectareas por personas') == clave and plano('Promedio') in columnas_base:
            destino_de[columna] = columnas_base[plano('Promedio')]
    columnas_pestana = {}
    for columna in range(1, ancho + 1):
        titulo = modelo.cell(row=FILA_CABECERAS_PESTANA, column=columna).value
        if titulo:
            columnas_pestana.setdefault(plano(titulo), columna)
    propias = {}
    for columna in range(1, ancho + 1):
        if columna in destino_de or not modelo.cell(row=FILA_CABECERAS_PESTANA, column=columna).value:
            continue
        for fila in range(primera, primera + FILAS):
            valor = modelo.cell(row=fila, column=columna).value
            if isinstance(valor, str) and valor.startswith('='):
                propias[columna] = (valor, fila)
                break

    for fila in range(primera, primera + FILAS):
        indice = ('=IFERROR(AGGREGATE(15,6,(ROW(%s)-%d)/(\'%s\'!$%s$%d:$%s$%d=$CB$%d),ROW()-%d),"")'
                  % (rango_filas, FILA_CABECERAS_BASE, HOJA_BASE, col_prov,
                     FILA_CABECERAS_BASE + 1, col_prov, 1005,
                     FILA_CABECERAS_PESTANA - 1, FILA_CABECERAS_PESTANA))
        lineas.append('  hoja.GetRange("CC%d").SetValue(%s);' % (fila, _js(indice)))
        for columna, columna_base in sorted(destino_de.items()):
            origen = "'%s'!%s$%d:%s$%d" % (HOJA_BASE, letra(columna_base),
                                           FILA_CABECERAS_BASE + 1,
                                           letra(columna_base), 1005)
            formula = ('=IF($CC%d="","",IF(INDEX(%s,$CC%d)="","",INDEX(%s,$CC%d)))'
                       % (fila, origen, fila, origen, fila))
            lineas.append('  hoja.GetRange("%s%d").SetValue(%s);'
                          % (letra(columna), fila, _js(formula)))
        for columna, (formula, fila_modelo) in sorted(propias.items()):
            # La fórmula de la plantilla, sin referencias de tabla y trasladada
            # a esta fila (la plantilla la trae de su primera fila con datos).
            trasladada = sin_referencias_de_tabla(formula, columnas_pestana, fila)
            if fila_modelo != fila:
                trasladada = re.sub(r'(?<![A-Z0-9$])([A-Z]{1,3})%d(?![0-9])' % fila_modelo,
                                    lambda m: '%s%d' % (m.group(1), fila), trasladada)
            cuerpo_formula = trasladada[1:]
            if not cuerpo_formula.upper().startswith(('IFERROR', 'SI.ERROR')):
                trasladada = '=IFERROR(%s,"")' % cuerpo_formula
            lineas.append('  hoja.GetRange("%s%d").SetValue(%s);'
                          % (letra(columna), fila, _js(trasladada)))
    # Fechas con su formato; si no, saldrían como número.
    for columna, columna_base in destino_de.items():
        celda = base.cell(row=FILA_CABECERAS_BASE + 1, column=columna_base)
        formato = str(celda.number_format or '')
        es_fecha_base = 'yy' in formato or 'mmm' in formato
        # La columna «Fecha» la escribe el sistema y siempre es una fecha,
        # tenga o no formato la celda que se mire de la hoja base.
        if es_fecha_base or plano(base.cell(row=FILA_CABECERAS_BASE,
                                            column=columna_base).value) in ('FECHA',):
            lineas.append('  hoja.GetRange("%s%d:%s%d").SetNumberFormat("%s");'
                          % (letra(columna), primera, letra(columna), primera + FILAS - 1,
                             'dd/mm/yyyy hh:mm' if columna == 1 else 'dd/mm/yyyy'))
    lineas.append('  hoja.GetRange("A%d:%s%d").SetBold(true);'
                  % (FILA_CABECERAS_PESTANA, letra(ancho), FILA_CABECERAS_PESTANA))
    lineas.append('}')
    return lineas, len(destino_de), len(propias)


def principal(usuario, ruta, aplicar):
    fisica = nucleo.ruta_fisica(usuario, ruta)
    cuerpo_guion, n_form, n_propias = guion(fisica)
    print('  %d columnas del formulario, %d propias' % (n_form, n_propias))
    token = firmar_jwt({'u': usuario, 'r': ruta, 'uso': 'descarga',
                        'exp': int(time.time()) + 900})
    url_libro = '%s/api/almacen/onlyoffice/download?t=%s' % (URL_PUBLICA, token)
    partes = (['builder.OpenFile("%s");' % url_libro] + cuerpo_guion +
              ['Api.RecalculateAllFormulas();',
               'builder.SaveFile("xlsx", "bolivar.xlsx");', 'builder.CloseFile();'])
    os.makedirs(CARPETA_SCRIPTS, exist_ok=True)
    fichero = secrets.token_hex(16) + '.docbuilder'
    camino = os.path.join(CARPETA_SCRIPTS, fichero)
    open(camino, 'w', encoding='utf-8').write('\n'.join(partes) + '\n')
    os.chmod(camino, 0o644)
    print('  guion de %d líneas' % len(partes))
    try:
        cuerpo = {'async': False, 'key': 'bol' + secrets.token_hex(8),
                  'url': '%s%s/%s' % (URL_PUBLICA, URL_SCRIPTS, fichero)}
        cuerpo['token'] = firmar_jwt(dict(cuerpo))
        respuesta = requests.post(
            url_interna_ds().rstrip('/') + '/docbuilder', json=cuerpo, timeout=(3, 900),
            headers={'Authorization': 'Bearer ' + firmar_jwt({'payload': cuerpo})})
        try:
            urls = list((respuesta.json() or {}).get('urls', {}).values())
        except ValueError:
            urls = []
        if not urls:
            print('  el docbuilder respondió:', respuesta.status_code, repr(respuesta.text[:300]))
            return 1
        salida = requests.get(_reescribir_url_interna(urls[0]), timeout=(3, 300)).content
        open('/tmp/bolivar.xlsx', 'wb').write(salida)
        print('  resultado en /tmp/bolivar.xlsx (%.0f KB)' % (len(salida) / 1024))
        if aplicar:
            import sala_editor
            from api_onlyoffice import _base_documento, invalidar_cache
            dentro = sala_editor.usuarios_conectados(_base_documento(usuario, ruta))
            if dentro is None or dentro:
                print('  NO se aplica: el Excel está abierto en el editor (%s)' % dentro)
                return 2
            carpeta, _, nombre = ruta.rpartition('/')
            nucleo.subir(usuario, carpeta or '/', nombre, io.BytesIO(salida))
            invalidar_cache(usuario, ruta)
            print('  aplicado al Drive')
    finally:
        os.remove(camino)
    return 0


if __name__ == '__main__':
    sys.exit(principal(int(sys.argv[1]), sys.argv[2], '--aplicar' in sys.argv))
