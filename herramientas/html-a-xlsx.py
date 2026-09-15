"""Convierte la exportación HTML de Google Sheets en UN libro .xlsx.

Google exporta cada hoja como una página HTML con sus estilos dentro (`.s0`,
`.s1`…). Aquí se leen esas páginas y se escribe un solo libro con una hoja por
página, conservando lo que se ve: texto, números, color de fondo, color y tipo
de letra, negrita y cursiva, alineación, celdas combinadas y los anchos y altos.
"""
import io
import os
import re
import sys

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def estilos_del_html(html):
    """El mapa clase → propiedades, leído del <style> de la propia página."""
    mapa = {}
    for bloque in re.findall(r'<style[^>]*>(.*?)</style>', html, re.S):
        for clase, cuerpo in re.findall(r'\.s(\d+)\s*\{([^}]*)\}', bloque):
            props = {}
            for trozo in cuerpo.split(';'):
                if ':' not in trozo:
                    continue
                nombre, valor = trozo.split(':', 1)
                props[nombre.strip()] = valor.strip()
            mapa['s' + clase] = props
    return mapa


def color(valor):
    """De `#rrggbb` o `rgb(r,g,b)` al formato de Excel (AARRGGBB)."""
    if not valor:
        return None
    valor = valor.strip()
    m = re.match(r'#([0-9a-fA-F]{6})$', valor)
    if m:
        return 'FF' + m.group(1).upper()
    m = re.match(r'#([0-9a-fA-F]{3})$', valor)
    if m:
        t = m.group(1)
        return 'FF' + ''.join(c * 2 for c in t).upper()
    m = re.match(r'rgb\((\d+),\s*(\d+),\s*(\d+)\)', valor)
    if m:
        return 'FF' + ''.join('%02X' % int(x) for x in m.groups())
    return None


def numero(texto):
    """Si la celda es un número (o un porcentaje), se guarda como número para
    que siga sirviendo en las fórmulas. Devuelve (valor, formato) o (None, None)."""
    t = (texto or '').strip()
    if not t or t in ('-', '—'):
        return None, None
    porciento = t.endswith('%')
    limpio = t[:-1].strip() if porciento else t
    moneda = limpio.startswith('$')
    if moneda:
        limpio = limpio[1:].strip()
    # miles con coma o punto, decimales con coma o punto
    if not re.match(r'^-?[\d.,]+$', limpio):
        return None, None
    solo = limpio
    if solo.count(',') and solo.count('.'):
        # el último separador manda como decimal
        if solo.rfind(',') > solo.rfind('.'):
            solo = solo.replace('.', '').replace(',', '.')
        else:
            solo = solo.replace(',', '')
    elif solo.count(','):
        partes = solo.split(',')
        solo = solo.replace(',', '.') if len(partes[-1]) != 3 else solo.replace(',', '')
    try:
        valor = float(solo)
    except ValueError:
        return None, None
    if porciento:
        return valor / 100.0, '0%' if valor == int(valor) else '0.00%'
    if moneda:
        return valor, '"$"#,##0.00'
    if valor == int(valor) and abs(valor) < 1e15:
        return int(valor), None
    return valor, None


def alto(estilo):
    m = re.search(r'height:\s*(\d+)', estilo or '')
    return int(m.group(1)) if m else None


def ancho(estilo):
    m = re.search(r'width:\s*(\d+)', estilo or '')
    return int(m.group(1)) if m else None


LADO_FINO = Side(style='thin', color='FFD0D0D0')


def pinta_celda(celda, props):
    """Le pone a la celda lo que dice el estilo de Google."""
    if not props:
        return
    fondo = color(props.get('background-color'))
    if fondo and fondo != 'FFFFFFFF':
        celda.fill = PatternFill('solid', fgColor=fondo)
    tam = props.get('font-size', '')
    m = re.match(r'([\d.]+)pt', tam)
    familia = (props.get('font-family') or '').split(',')[0].replace('docs-', '').strip()
    celda.font = Font(
        name=familia or 'Calibri',
        size=float(m.group(1)) if m else 11,
        bold=props.get('font-weight') in ('bold', '700'),
        italic=props.get('font-style') == 'italic',
        underline='single' if 'underline' in (props.get('text-decoration') or '') else None,
        color=color(props.get('color')) or 'FF000000')
    horizontal = props.get('text-align')
    vertical = props.get('vertical-align')
    celda.alignment = Alignment(
        horizontal=horizontal if horizontal in ('left', 'center', 'right', 'justify') else None,
        vertical={'top': 'top', 'middle': 'center', 'bottom': 'bottom'}.get(vertical),
        wrap_text=(props.get('white-space') == 'normal'))
    if props.get('border-bottom') or props.get('border'):
        celda.border = Border(bottom=LADO_FINO)


def pasa_una_hoja(libro, ruta, nombre):
    html = open(ruta, encoding='utf-8', errors='replace').read()
    estilos = estilos_del_html(html)
    sopa = BeautifulSoup(html, 'html.parser')
    tabla = sopa.find('table', class_='waffle') or sopa.find('table')
    if tabla is None:
        return None
    hoja = libro.create_sheet(nombre[:31])

    # anchos de columna, de las cabeceras de la tabla
    cabecera = tabla.find('thead')
    if cabecera:
        columnas = [th for th in cabecera.find_all('th')
                    if 'column-headers-background' in (th.get('class') or [])]
        for i, th in enumerate(columnas, start=1):
            px = ancho(th.get('style'))
            if px:
                hoja.column_dimensions[get_column_letter(i)].width = max(2.0, px / 7.0)

    ocupadas = {}          # (fila, col) reservadas por celdas combinadas
    combinar = []
    fila_n = 0
    for tr in tabla.find_all('tr'):
        celdas = tr.find_all(['td', 'th'], recursive=False)
        # la primera celda de cada fila es el número de fila: no es dato
        if celdas and 'row-headers-background' in (celdas[0].get('class') or []):
            celdas = celdas[1:]
        elif celdas and celdas[0].name == 'th':
            continue                      # la fila de cabecera de columnas
        fila_n += 1
        px = alto(tr.get('style'))
        if px:
            hoja.row_dimensions[fila_n].height = px * 0.75
        col_n = 0
        for celda in celdas:
            col_n += 1
            while ocupadas.get((fila_n, col_n)):
                col_n += 1
            texto = celda.get_text(' ', strip=True)
            destino = hoja.cell(row=fila_n, column=col_n)
            if texto:
                valor, formato = numero(texto)
                if valor is None:
                    destino.value = texto
                else:
                    destino.value = valor
                    if formato:
                        destino.number_format = formato
            clases = [c for c in (celda.get('class') or []) if re.match(r'^s\d+$', c)]
            pinta_celda(destino, estilos.get(clases[0]) if clases else None)

            filas = int(celda.get('rowspan') or 1)
            cols = int(celda.get('colspan') or 1)
            if filas > 1 or cols > 1:
                combinar.append((fila_n, col_n, fila_n + filas - 1, col_n + cols - 1))
                for f in range(fila_n, fila_n + filas):
                    for c in range(col_n, col_n + cols):
                        ocupadas[(f, c)] = True
                col_n += cols - 1

    for f1, c1, f2, c2 in combinar:
        try:
            hoja.merge_cells(start_row=f1, start_column=c1, end_row=f2, end_column=c2)
        except Exception:
            pass
    hoja.sheet_view.showGridLines = False
    return hoja


def es_exportacion_de_google(carpeta):
    """¿Esta carpeta es la exportación HTML de un libro de Google Sheets?

    Google escribe una página por hoja y las marca con sus clases `.ritz .waffle`.
    Basta con encontrar esa firma en cualquiera de las páginas."""
    try:
        paginas = [f for f in os.listdir(carpeta) if f.lower().endswith('.html')]
    except OSError:
        return False
    for pagina in paginas:
        try:
            with io.open(os.path.join(carpeta, pagina), encoding='utf-8',
                         errors='ignore') as fichero:
                if 'waffle' in fichero.read(20000):
                    return True
        except OSError:
            continue
    return False


def convertir(carpeta, salida):
    """Escribe un `.xlsx` con una hoja por página HTML. Devuelve cuántas hojas
    salieron (0 si la carpeta no traía páginas)."""
    paginas = sorted(f for f in os.listdir(carpeta) if f.lower().endswith('.html'))
    if not paginas:
        return 0
    libro = Workbook()
    libro.remove(libro.active)
    hechas = 0
    for pagina in paginas:
        nombre = os.path.splitext(pagina)[0]
        if pasa_una_hoja(libro, os.path.join(carpeta, pagina), nombre) is not None:
            hechas += 1
    if not hechas:
        return 0
    libro.save(salida)
    return hechas


def main(carpeta, salida):
    paginas = sorted(f for f in os.listdir(carpeta) if f.lower().endswith('.html'))
    if not paginas:
        print('no hay páginas HTML en', carpeta)
        return 1
    libro = Workbook()
    libro.remove(libro.active)
    for pagina in paginas:
        nombre = os.path.splitext(pagina)[0]
        hoja = pasa_una_hoja(libro, os.path.join(carpeta, pagina), nombre)
        print('  hoja «%s»: %s' % (nombre, 'sin tabla' if hoja is None
                                   else '%d filas × %d columnas' % (hoja.max_row, hoja.max_column)))
    libro.save(salida)
    print('escrito:', salida, '(%.1f KB)' % (os.path.getsize(salida) / 1024.0))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1], sys.argv[2]))
